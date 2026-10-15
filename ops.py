"""
운영 확장 기능 (데이터 계층)

예산 · 공급업체 · 공용 장비 대여 · 첨부파일 · 점검/법정검사 · 보증 · 직원별 자산과 퇴사 회수 ·
소프트웨어 라이선스 · 렌탈/리스 · 일괄 수정 · 자본적 지출 · 회계 분개 · 분석 · 알림 · 법인차량 운행 연동
화면 코드는 이 모듈과 data.py의 함수만 호출합니다.
"""
from __future__ import annotations

import json
import re
import smtplib
from datetime import date, timedelta
from email.mime.text import MIMEText
from pathlib import Path

import data as d
from data import DataError, clean_str, to_date, to_num

try:
    import requests
except ImportError:  # requirements 설치 전
    requests = None

# 역할 이름(auth.py와 같음 — 데이터 계층은 화면 모듈을 불러오지 않습니다)
ADMIN, EXEC, FIN, LEAD, SITE, STAFF = "관리자", "임원", "재무", "현장책임자", "현장담당", "직원"
SITE_ROLES = (LEAD, SITE)


def _new(table: str, fields: dict) -> dict:
    rec = {c: "" for c in d.TABLES[table]["cols"]}
    rec.update({k: v for k, v in fields.items() if k in rec})
    return rec


def _log(rec: dict, action: str, actor: str, note: str = ""):
    rec["처리이력"] = (rec.get("처리이력") + "\n" if rec.get("처리이력") else "") + \
        f"[{d.now_text()}] {action} · {actor}" + (f" · {note}" if note else "")


# =========================================================
# 예산
# =========================================================
def get_budget(year: int, site: str) -> float:
    b = d.lookup("budgets", (int(year), site))
    return float(b["예산액"] or 0) if b else 0.0


def save_budgets(year: int, amounts: dict[str, float], actor: str):
    rows = []
    for site, amt in amounts.items():
        old = d.lookup("budgets", (int(year), site)) or {}
        rows.append({"연도": int(year), "현장": site, "예산액": to_num(amt, 0) or 0, "비고": old.get("비고", "")})
    d.upsert("budgets", rows)
    d.add_history("-", "예산 저장", f"{year}년 " + ", ".join(f"{s} {int(to_num(a, 0) or 0):,}원" for s, a in amounts.items()), actor)


def budget_status(year: int, site: str | None = None) -> list[dict]:
    """현장별 예산 대비 집행(발주·등록 완료)과 진행 중(결재 중) 금액."""
    out = []
    sites = [site] if site else d.site_names()
    for s in sites:
        used = pending = 0.0
        for r in d.load("pr"):
            if r["현장"] != s or not r["요청일시"].startswith(str(year)) or r["상태"] == d.REJECTED:
                continue
            if r["상태"] in ("발주", "검수완료", "등록완료"):
                used += d.pr_amount(r)
            else:
                pending += d.pr_amount(r)
        budget = get_budget(year, s)
        out.append({"현장": s, "예산": budget, "집행": used, "결재 중": pending, "잔액": budget - used - pending,
                    "집행률(%)": round(used / budget * 100, 1) if budget else None})
    return out


def budget_check(site: str, amount: float, year: int | None = None, exclude_pr: str | None = None) -> dict:
    """구매 요청 금액을 반영했을 때 예산이 남는지."""
    year = year or d.today().year
    st = budget_status(year, site)[0]
    if exclude_pr:
        rec = d.get_request("pr", exclude_pr)
        if rec and rec["현장"] == site and rec["요청일시"].startswith(str(year)):
            st["결재 중"] -= d.pr_amount(rec) if rec["상태"] in ("요청", "승인") else 0
            st["잔액"] = st["예산"] - st["집행"] - st["결재 중"]
    after = st["잔액"] - amount
    return {**st, "요청 금액": amount, "요청 후 잔액": after, "초과": st["예산"] > 0 and after < 0, "예산 없음": st["예산"] <= 0}


# =========================================================
# 공급업체
# =========================================================
def vendor_names(active_only: bool = True) -> list[str]:
    return sorted(v["업체명"] for v in d.load("vendors") if not active_only or v["상태"] != "거래중지")


def save_vendor(fields: dict, actor: str, original: str | None = None):
    name = clean_str(fields.get("업체명"))
    if not name:
        raise DataError("업체명을 입력하세요.")
    if original and original != name:
        raise DataError("업체명은 바꿀 수 없습니다. 새 업체로 등록하고 기존 업체는 '거래중지'로 바꾸세요.")
    if not original and d.lookup("vendors", name):
        raise DataError(f"'{name}' 업체가 이미 있습니다.")
    old = d.lookup("vendors", name) or {}
    rec = _new("vendors", {**old, **fields})
    rec["업체명"] = name
    rec["상태"] = rec["상태"] or "거래중"
    rec["등록일시"] = old.get("등록일시") or d.now_text()
    d.upsert("vendors", [rec])


def vendor_stats() -> list[dict]:
    """업체별 견적 참여·선정·거래금액·평균 납기·검수 불합격·현재 보유 자산."""
    stats: dict[str, dict] = {}

    def s(name):
        return stats.setdefault(name, {"업체명": name, "견적 참여": 0, "선정": 0, "거래금액": 0.0, "납기 합계": 0,
                                       "납기 건수": 0, "검수 불합격": 0, "보유 자산": 0})
    for r in d.load("pr"):
        for q in d.pr_quotes(r):
            s(q["업체"])["견적 참여"] += 1
        if r["선정업체"]:
            v = s(r["선정업체"])
            v["선정"] += 1
            v["거래금액"] += (r["선정단가"] or 0) * (r["입고누계"] or r["수량"] or 0)
            v["검수 불합격"] += r["처리이력"].count("검수 부적합")
            if r["발주일시"] and r["검수일시"]:
                days = (to_date(r["검수일시"]) - to_date(r["발주일시"])).days
                v["납기 합계"] += max(days, 0)
                v["납기 건수"] += 1
    for a in d.active_assets():
        if a["공급업체"]:
            s(a["공급업체"])["보유 자산"] += 1
    known = {v["업체명"]: v for v in d.load("vendors")}
    out = []
    for name, v in stats.items():
        v["평균 납기(일)"] = round(v["납기 합계"] / v["납기 건수"], 1) if v["납기 건수"] else None
        v["등록"] = "등록" if name in known else "미등록"
        v["상태"] = known.get(name, {}).get("상태", "")
        out.append({k: v[k] for k in ("업체명", "등록", "상태", "견적 참여", "선정", "거래금액", "평균 납기(일)", "검수 불합격", "보유 자산")})
    for name, v in known.items():
        if name not in stats:
            out.append({"업체명": name, "등록": "등록", "상태": v["상태"], "견적 참여": 0, "선정": 0, "거래금액": 0,
                        "평균 납기(일)": None, "검수 불합격": 0, "보유 자산": 0})
    return sorted(out, key=lambda x: (-x["선정"], -x["견적 참여"], x["업체명"]))


# =========================================================
# 공용 장비 대여
# =========================================================
RENTAL_FLOW = ["신청", "승인", "대여중", "반납완료"]
RENTAL_OPEN = ("신청", "승인", "대여중")
MAX_RENTAL_DAYS = 30


def rentable_assets(site: str | None = None) -> list[dict]:
    return [a for a in d.active_assets() if a["공용"] == "Y" and a["상태"] in ("사용중", "유휴") and (not site or a["현장"] == site)]


def get_rentals(site: str | None = None) -> list[dict]:
    rows = [r for r in d.load("rentals") if not site or r["현장"] == site]
    return sorted(rows, key=lambda r: (r["시작일"] or date.min, r["대여번호"]), reverse=True)


def rental_conflict(tag: str, start: date, end: date, exclude: str | None = None) -> dict | None:
    for r in d.load("rentals"):
        if r["태그번호"] != tag or r["상태"] not in RENTAL_OPEN or r["대여번호"] == exclude:
            continue
        if r["시작일"] and r["반납예정일"] and start <= r["반납예정일"] and r["시작일"] <= end:
            return r
    return None


def create_rental(tag: str, start, end, purpose: str, user_id: str, user_name: str, dept: str, actor: str) -> str:
    a = d.get_asset(tag)
    start, end = to_date(start), to_date(end)
    if not a or a["공용"] != "Y" or a["상태"] not in ("사용중", "유휴"):
        raise DataError("대여할 수 있는 공용 장비가 아닙니다.")
    if not start or not end or end < start:
        raise DataError("대여 기간을 확인하세요 (반납일이 시작일보다 빠릅니다).")
    if start < d.today():
        raise DataError("지난 날짜로는 신청할 수 없습니다.")
    if (end - start).days + 1 > MAX_RENTAL_DAYS:
        raise DataError(f"한 번에 {MAX_RENTAL_DAYS}일까지 빌릴 수 있습니다.")
    if not clean_str(purpose):
        raise DataError("사용 목적을 입력하세요.")
    with d._lock:
        c = rental_conflict(tag, start, end)
        if c:
            raise DataError(f"이미 예약된 기간입니다: {c['시작일']} ~ {c['반납예정일']} ({c['신청자명']}, {c['상태']})")
        no = d.next_code("rentals", "대여번호", "RT")
        rec = _new("rentals", {"대여번호": no, "상태": "신청", "태그번호": tag, "품명": a["품명"], "현장": a["현장"],
                               "신청자": user_id, "신청자명": user_name, "부서": dept, "시작일": start, "반납예정일": end,
                               "사용목적": purpose, "신청일시": d.now_text()})
        _log(rec, "대여 신청", actor, purpose)
        d.upsert("rentals", [rec])
    return no


def _update_rental(no: str, expect: tuple, mutate):
    with d._lock:
        cur = d.lookup("rentals", no)
        if not cur:
            raise DataError("대여 신청을 찾을 수 없습니다.")
        if cur["상태"] not in expect:
            raise DataError(f"{no}은(는) 이미 '{cur['상태']}' 상태입니다.")
        rec = dict(cur)
        mutate(rec)
        d.upsert("rentals", [rec])
        return rec


def rental_decide(no: str, approve: bool, note: str, actor: str):
    if not approve and not clean_str(note):
        raise DataError("반려 사유를 입력하세요.")

    def m(rec):
        rec["상태"] = "승인" if approve else d.REJECTED
        rec["승인자"], rec["승인일시"] = actor, d.now_text()
        _log(rec, "승인" if approve else "반려", actor, note)
    _update_rental(no, ("신청",), m)


def rental_handout(no: str, actor: str):
    def m(rec):
        rec["상태"], rec["대여일시"] = "대여중", d.now_text()
        _log(rec, "장비 인계(대여 시작)", actor)
    rec = _update_rental(no, ("승인",), m)
    d.add_history(rec["태그번호"], "대여", f"{rec['신청자명']} · {rec['시작일']} ~ {rec['반납예정일']} ({no})", actor)


def rental_return(no: str, condition: str, note: str, actor: str) -> dict:
    def m(rec):
        rec["상태"], rec["반납일시"], rec["반납상태"], rec["메모"] = "반납완료", d.now_text(), condition, note
        late = (d.today() - rec["반납예정일"]).days if rec["반납예정일"] else 0
        _log(rec, "반납" + (f" (연체 {late}일)" if late > 0 else ""), actor, f"{condition} {note}".strip())
    rec = _update_rental(no, ("대여중",), m)
    d.add_history(rec["태그번호"], "반납", f"{rec['신청자명']} · 상태 {condition}" + (f" · {note}" if note else "") + f" ({no})", actor)
    return rec


def rental_cancel(no: str, actor: str):
    def m(rec):
        rec["상태"] = "취소"
        _log(rec, "신청 취소", actor)
    _update_rental(no, ("신청", "승인"), m)


def overdue_rentals(site: str | None = None) -> list[dict]:
    t = d.today()
    return [r for r in d.load("rentals") if r["상태"] == "대여중" and r["반납예정일"] and r["반납예정일"] < t
            and (not site or r["현장"] == site)]


# =========================================================
# 첨부파일
# =========================================================
ATTACH_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".pdf", ".xlsx", ".xls", ".docx", ".hwp", ".hwpx", ".txt", ".csv", ".zip"}
ATTACH_MAX = 10 * 1024 * 1024
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp"}


def attach_dir() -> Path:
    p = d.DATA_DIR / "attachments"
    p.mkdir(parents=True, exist_ok=True)
    return p


def save_attachment(target: str, target_no: str, filename: str, blob: bytes, desc: str, actor: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ATTACH_EXT:
        raise DataError(f"{ext or '확장자 없는'} 파일은 올릴 수 없습니다. 사진·PDF·문서 파일을 올려 주세요.")
    if len(blob) > ATTACH_MAX:
        raise DataError("파일은 10MB까지 올릴 수 있습니다.")
    with d._lock:
        no = d.next_code_query("attachments", "첨부번호", "AT")
        safe = re.sub(r"[^\w.\-가-힣]", "_", Path(filename).name)[-80:]
        stored = f"{no}_{safe}"
        (attach_dir() / stored).write_bytes(blob)
        d.insert_bulk("attachments", [{"첨부번호": no, "대상": target, "대상번호": target_no, "파일명": Path(filename).name,
                                       "저장명": stored, "형식": ext.lstrip("."), "크기": len(blob), "설명": desc,
                                       "등록자": actor, "등록일시": d.now_text()}])
    if target == "asset":
        d.add_history(target_no, "첨부", f"{Path(filename).name}" + (f" · {desc}" if desc else ""), actor)
    return no


def list_attachments(target: str, target_no: str) -> list[dict]:
    return d.query("attachments", '"대상"=? AND "대상번호"=?', (target, target_no), order='"등록일시" DESC')


def read_attachment(rec: dict) -> bytes | None:
    p = attach_dir() / rec["저장명"]
    return p.read_bytes() if p.exists() else None


def delete_attachment(no: str, actor: str):
    rows = d.query("attachments", '"첨부번호"=?', (no,))
    if not rows:
        return
    p = attach_dir() / rows[0]["저장명"]
    if p.exists():
        p.unlink()
    d.delete_where("attachments", '"첨부번호"=?', (no,))
    if rows[0]["대상"] == "asset":
        d.add_history(rows[0]["대상번호"], "첨부 삭제", rows[0]["파일명"], actor)


def is_image(rec: dict) -> bool:
    return "." + rec["형식"].lower() in IMAGE_EXT


# =========================================================
# 점검 · 법정검사 · 보증
# =========================================================
MAINT_KINDS = ["정기점검", "법정검사", "검교정", "수리", "소모품교체"]


def save_plan(fields: dict, actor: str) -> str:
    tag = clean_str(fields.get("태그번호"))
    if not d.get_asset(tag):
        raise DataError("자산을 찾을 수 없습니다.")
    if not clean_str(fields.get("점검명")):
        raise DataError("점검명을 입력하세요.")
    months = int(to_num(fields.get("주기개월"), 0) or 0)
    if months <= 0:
        raise DataError("점검 주기(개월)를 입력하세요.")
    with d._lock:
        no = clean_str(fields.get("계획번호")) or d.next_code("maint_plans", "계획번호", "MP")
        old = d.lookup("maint_plans", no) or {}
        rec = _new("maint_plans", {**old, **fields, "계획번호": no})
        if not rec["다음점검일"]:
            base = to_date(rec["최근점검일"]) or d.today()
            rec["다음점검일"] = d._shift_months_exact(base, months) if rec["최근점검일"] else base
        rec["활성"] = rec["활성"] or "Y"
        d.upsert("maint_plans", [rec])
    if not old:
        d.add_history(tag, "점검 계획", f"{rec['점검명']} · {months}개월마다 · 다음 {rec['다음점검일']}", actor)
    return no


def plans_for(tag: str) -> list[dict]:
    return [p for p in d.load("maint_plans") if p["태그번호"] == tag and p["활성"] == "Y"]


def logs_for(tag: str) -> list[dict]:
    return sorted([r for r in d.load("maint_logs") if r["태그번호"] == tag], key=lambda r: r["일자"] or date.min, reverse=True)


def plans_due(days: int | None = None, site: str | None = None) -> list[dict]:
    """다음 점검일이 지났거나 days일 안에 오는 계획(보유 자산만)."""
    days = d.setting_num("알림 기준일", 30) if days is None else days
    limit = d.today() + timedelta(days=days)
    out = []
    for p in d.load("maint_plans"):
        if p["활성"] != "Y" or not p["다음점검일"] or p["다음점검일"] > limit:
            continue
        a = d.get_asset(p["태그번호"])
        if not a or not d.is_active(a) or (site and a["현장"] != site):
            continue
        left = (p["다음점검일"] - d.today()).days
        out.append({**p, "품명": a["품명"], "현장": a["현장"], "남은 일수": left, "상태": "지연" if left < 0 else "임박"})
    return sorted(out, key=lambda r: r["다음점검일"])


def warranty_expiring(days: int | None = None, site: str | None = None) -> list[dict]:
    days = d.setting_num("알림 기준일", 30) if days is None else days
    t = d.today()
    return sorted([a for a in d.active_assets() if a["보증만료일"] and t <= a["보증만료일"] <= t + timedelta(days=days)
                   and (not site or a["현장"] == site)], key=lambda a: a["보증만료일"])


# =========================================================
# 직원 · 퇴사 회수
# =========================================================
def get_employees(site: str | None = None, status: str | None = None) -> list[dict]:
    rows = [e for e in d.load("employees") if (not site or e["현장"] == site) and (not status or e["상태"] == status)]
    return sorted(rows, key=lambda e: (e["현장"], e["부서"], e["이름"]))


def save_employee(fields: dict, actor: str) -> str:
    if not clean_str(fields.get("이름")):
        raise DataError("이름을 입력하세요.")
    if clean_str(fields.get("현장")) not in d.site_names():
        raise DataError("소속 현장을 고르세요.")
    with d._lock:
        no = clean_str(fields.get("사번")) or d.next_code("employees", "사번", "E")
        old = d.lookup("employees", no) or {}
        rec = _new("employees", {**old, **fields, "사번": no})
        rec["상태"] = rec["상태"] or "재직"
        d.upsert("employees", [rec])
    return no


def import_employees(rows: list[dict], actor: str) -> tuple[int, list[str]]:
    sites, errors, recs = set(d.site_names()), [], []
    for i, r in enumerate(rows, start=2):
        name, site = clean_str(r.get("이름")), clean_str(r.get("현장"))
        if not name:
            continue
        if site not in sites:
            errors.append(f"{i}행 {name}: 현장 '{site}' 없음")
            continue
        recs.append(r)
    if errors:
        return 0, errors
    for r in recs:
        save_employee({k: r.get(k) for k in ("사번", "이름", "현장", "부서", "직급", "이메일", "연락처", "입사일")}, actor)
    return len(recs), []


def employee_assets(emp: dict) -> list[dict]:
    return [a for a in d.active_assets() if (a["사용자사번"] and a["사용자사번"] == emp["사번"]) or
            (not a["사용자사번"] and a["사용자"] == emp["이름"] and a["현장"] == emp["현장"])]


def employee_licenses(emp: dict) -> list[dict]:
    return [x for x in d.load("license_assign") if x["상태"] == "사용중" and
            ((x["사번"] and x["사번"] == emp["사번"]) or (not x["사번"] and x["사용자"] == emp["이름"]))]


def holders_summary(site: str | None = None) -> list[dict]:
    """직원별 보유 자산 수·금액."""
    emps = get_employees(site)
    by_no = {e["사번"]: e for e in emps}
    by_name = {(e["이름"], e["현장"]): e for e in emps}
    agg: dict[str, dict] = {}
    for a in d.active_assets():
        if site and a["현장"] != site:
            continue
        e = by_no.get(a["사용자사번"]) if a["사용자사번"] else by_name.get((a["사용자"], a["현장"]))
        if not e:
            continue
        g = agg.setdefault(e["사번"], {"보유 자산": 0, "취득가액": 0.0})
        g["보유 자산"] += 1
        g["취득가액"] += a["취득가액"] or 0
    return [{**e, **agg.get(e["사번"], {"보유 자산": 0, "취득가액": 0})} for e in emps]


def start_return(emp_no: str, retire_date, actor: str) -> str:
    """퇴사 회수 시작: 보유 자산 체크리스트를 만들고 직원 상태를 '퇴사예정'으로 바꿉니다."""
    emp = d.lookup("employees", emp_no)
    if not emp:
        raise DataError("직원을 찾을 수 없습니다.")
    if any(r["사번"] == emp_no and r["상태"] == "진행중" for r in d.load("returns")):
        raise DataError("이미 진행 중인 회수 건이 있습니다.")
    assets = employee_assets(emp)
    with d._lock:
        no = d.next_code("returns", "회수번호", "RC")
        d.upsert("returns", [{"회수번호": no, "사번": emp_no, "이름": emp["이름"], "현장": emp["현장"], "부서": emp["부서"],
                              "퇴사일": to_date(retire_date) or d.today(), "상태": "진행중" if assets else "완료",
                              "생성자": actor, "생성일시": d.now_text(), "완료일시": "" if assets else d.now_text()}])
        d.insert_bulk("return_items", [{"회수번호": no, "태그번호": a["태그번호"], "품명": a["품명"], "결과": "미처리",
                                        "인계자": "", "메모": "", "처리자": "", "처리일시": ""} for a in assets])
        d.upsert("employees", [dict(emp, 상태="퇴사예정" if assets else "퇴사", 퇴사일=to_date(retire_date) or d.today())])
    for a in assets:
        d.add_history(a["태그번호"], "퇴사 회수", f"{emp['이름']} 퇴사 예정 ({no})", actor)
    if not assets:
        _release_licenses(emp, actor)
    return no


def return_items(no: str) -> list[dict]:
    return d.query("return_items", '"회수번호"=?', (no,), order='"태그번호"')


def process_return_item(no: str, tag: str, result: str, handover_to: str, memo: str, actor: str) -> bool:
    """회수 항목 처리. result: 회수(→유휴) / 인계(→새 사용자) / 분실(→경위서 요청). 모두 끝나면 True."""
    if result not in ("회수", "인계", "분실"):
        raise DataError("처리 결과를 고르세요.")
    if result == "인계" and not clean_str(handover_to):
        raise DataError("인계받을 사람을 입력하세요.")
    ret = d.lookup("returns", no)
    if not ret or ret["상태"] != "진행중":
        raise DataError("진행 중인 회수 건이 아닙니다.")
    items = {r["태그번호"]: r for r in return_items(no)}
    if tag not in items:
        raise DataError("회수 목록에 없는 자산입니다.")
    a = d.get_asset(tag)
    if result == "회수" and a:
        d.update_asset(tag, {"사용자": "", "사용자사번": "", "상태": "유휴" if a["상태"] == "사용중" else a["상태"]},
                       actor, "퇴사 회수", f"{ret['이름']} 퇴사 → 회수 (재배치 대상) ({no})")
    elif result == "인계" and a:
        emp = next((e for e in get_employees(ret["현장"]) if e["이름"] == clean_str(handover_to) and e["상태"] == "재직"), None)
        d.update_asset(tag, {"사용자": clean_str(handover_to), "사용자사번": emp["사번"] if emp else ""},
                       actor, "퇴사 인계", f"{ret['이름']} → {handover_to} ({no})")
    elif result == "분실":
        d.request_statement(tag, f"퇴사자 {ret['이름']} 미반납", actor)
    item = dict(items[tag], 결과=result, 인계자=clean_str(handover_to), 메모=memo, 처리자=actor, 처리일시=d.now_text())
    d.insert_bulk("return_items", [item], normalized=True)
    left = [r for r in return_items(no) if r["결과"] == "미처리"]
    if not left:
        d.upsert("returns", [dict(ret, 상태="완료", 완료일시=d.now_text())])
        emp = d.lookup("employees", ret["사번"])
        if emp:
            d.upsert("employees", [dict(emp, 상태="퇴사")])
            _release_licenses(emp, actor)
        return True
    return False


def _release_licenses(emp: dict, actor: str):
    for x in employee_licenses(emp):
        release_license(x["할당번호"], actor, reason=f"{emp['이름']} 퇴사")


def get_returns(site: str | None = None) -> list[dict]:
    return sorted([r for r in d.load("returns") if not site or r["현장"] == site], key=lambda r: r["생성일시"], reverse=True)


# =========================================================
# 소프트웨어 라이선스
# =========================================================
def save_license(fields: dict, actor: str) -> str:
    if not clean_str(fields.get("소프트웨어명")):
        raise DataError("소프트웨어명을 입력하세요.")
    if (to_num(fields.get("보유수량"), 0) or 0) < 1:
        raise DataError("보유 수량을 1 이상으로 입력하세요.")
    with d._lock:
        no = clean_str(fields.get("라이선스번호")) or d.next_code("licenses", "라이선스번호", "LC")
        old = d.lookup("licenses", no) or {}
        if old and (to_num(fields.get("보유수량"), 0) or 0) < license_used(no):
            raise DataError(f"이미 {license_used(no)}명에게 할당되어 있어 보유 수량을 그보다 줄일 수 없습니다.")
        rec = _new("licenses", {**old, **fields, "라이선스번호": no})
        rec["상태"] = rec["상태"] or "사용"
        d.upsert("licenses", [rec])
    return no


def license_used(no: str) -> int:
    return sum(1 for x in d.load("license_assign") if x["라이선스번호"] == no and x["상태"] == "사용중")


def assign_license(no: str, user: str, emp_no: str, dept: str, site: str, pc_tag: str, actor: str) -> str:
    lic = d.lookup("licenses", no)
    if not lic or lic["상태"] != "사용":
        raise DataError("사용 중인 라이선스가 아닙니다.")
    if not clean_str(user):
        raise DataError("사용자를 입력하세요.")
    if license_used(no) >= int(lic["보유수량"] or 0):
        raise DataError(f"남은 라이선스가 없습니다 (보유 {int(lic['보유수량'])}개 모두 할당). 수량을 늘리거나 다른 사용자를 해제하세요.")
    if any(x["라이선스번호"] == no and x["상태"] == "사용중" and x["사용자"] == clean_str(user) and x["사번"] == clean_str(emp_no)
           for x in d.load("license_assign")):
        raise DataError(f"{user}님에게 이미 할당되어 있습니다.")
    with d._lock:
        aid = d.next_code("license_assign", "할당번호", "LA")
        d.upsert("license_assign", [{"할당번호": aid, "라이선스번호": no, "사용자": clean_str(user), "사번": clean_str(emp_no),
                                     "부서": dept, "현장": site, "PC태그": clean_str(pc_tag).upper(), "할당일": d.today(),
                                     "해제일": None, "상태": "사용중"}])
    if pc_tag and d.get_asset(pc_tag):
        d.add_history(pc_tag.upper(), "라이선스 설치", f"{lic['소프트웨어명']} · {user}", actor)
    return aid


def release_license(aid: str, actor: str, reason: str = ""):
    x = d.lookup("license_assign", aid)
    if not x or x["상태"] != "사용중":
        return
    d.upsert("license_assign", [dict(x, 상태="해제", 해제일=d.today())])
    if x["PC태그"] and d.get_asset(x["PC태그"]):
        lic = d.lookup("licenses", x["라이선스번호"]) or {}
        d.add_history(x["PC태그"], "라이선스 해제", f"{lic.get('소프트웨어명', '')} · {x['사용자']} {reason}".strip(), actor)


def license_usage() -> list[dict]:
    t = d.today()
    out = []
    for lic in sorted(d.load("licenses"), key=lambda r: r["소프트웨어명"]):
        used = license_used(lic["라이선스번호"])
        total = int(lic["보유수량"] or 0)
        left = (lic["만료일"] - t).days if lic["만료일"] else None
        out.append({**lic, "사용": used, "남음": total - used, "사용률(%)": round(used / total * 100, 1) if total else 0,
                    "만료까지(일)": left})
    return out


def licenses_expiring(days: int | None = None) -> list[dict]:
    days = d.setting_num("알림 기준일", 30) * 2 if days is None else days
    return [x for x in license_usage() if x["상태"] == "사용" and x["만료까지(일)"] is not None and x["만료까지(일)"] <= days]


# =========================================================
# 렌탈 · 리스 (회사 소유가 아닌 장비)
# =========================================================
def save_lease(fields: dict, actor: str) -> str:
    if not clean_str(fields.get("품명")):
        raise DataError("품명을 입력하세요.")
    if clean_str(fields.get("현장")) not in d.site_names():
        raise DataError("현장을 고르세요.")
    s, e = to_date(fields.get("시작일")), to_date(fields.get("종료일"))
    if s and e and e < s:
        raise DataError("종료일이 시작일보다 빠릅니다.")
    with d._lock:
        no = clean_str(fields.get("계약번호")) or d.next_code("leases", "계약번호", "LS")
        old = d.lookup("leases", no) or {}
        rec = _new("leases", {**old, **fields, "계약번호": no})
        rec["상태"] = rec["상태"] or "사용중"
        d.upsert("leases", [rec])
    return no


def get_leases(site: str | None = None, active_only: bool = False) -> list[dict]:
    rows = [x for x in d.load("leases") if (not site or x["현장"] == site) and (not active_only or x["상태"] == "사용중")]
    return sorted(rows, key=lambda x: (x["종료일"] or date.max, x["계약번호"]))


def leases_expiring(days: int | None = None, site: str | None = None) -> list[dict]:
    days = d.setting_num("알림 기준일", 30) * 2 if days is None else days
    t = d.today()
    return [x for x in get_leases(site, True) if x["종료일"] and x["종료일"] <= t + timedelta(days=days)]


def lease_monthly_cost(site: str | None = None) -> float:
    return sum(x["월요금"] or 0 for x in get_leases(site, True))


# =========================================================
# 일괄 수정 · 자본적 지출
# =========================================================
BULK_FIELDS = ["관리부서", "사용자", "사용자사번", "설치위치", "상태", "공용", "보증만료일", "세무 내용연수", "세무 상각방법", "비고"]


def bulk_update(tags: list[str], changes: dict, actor: str) -> int:
    changes = {k: v for k, v in changes.items() if k in BULK_FIELDS}
    if not tags or not changes:
        raise DataError("바꿀 자산과 항목을 고르세요.")
    if "상태" in changes and changes["상태"] not in ("사용중", "유휴"):
        raise DataError("일괄 수정으로는 '사용중'이나 '유휴'로만 바꿀 수 있습니다.")
    by_tag, texts = {}, {}
    for t in tags:
        a = d.get_asset(t)
        if not a or not d.is_active(a):
            continue
        ch = dict(changes)
        if "상태" in ch and a["상태"] not in ("사용중", "유휴"):
            ch.pop("상태")
        by_tag[t] = ch
        texts[t] = "일괄 수정: " + ", ".join(f"{k} {clean_str(a[k]) or '-'} → {clean_str(v) or '-'}" for k, v in ch.items())
    if not by_tag:
        raise DataError("수정할 수 있는 자산이 없습니다.")
    d.update_assets_bulk(by_tag, actor, "일괄 수정", texts)
    return len(by_tag)


def add_capex(tag: str, when, content: str, amount, extend_years, actor: str) -> str:
    a = d.get_asset(tag)
    when = to_date(when)
    if not a or not d.is_active(a):
        raise DataError("보유 중인 자산에만 자본적 지출을 더할 수 있습니다.")
    if d.account_info(a)["cls"] == "소모품":
        raise DataError("소모품에는 자본적 지출을 더할 수 없습니다.")
    if not when or (a["취득일자"] and when < a["취득일자"]):
        raise DataError("지출일은 취득일 이후여야 합니다.")
    if (to_num(amount, 0) or 0) <= 0:
        raise DataError("금액을 입력하세요.")
    if not clean_str(content):
        raise DataError("내용을 입력하세요 (예: 엔진 오버홀, 용량 증설).")
    with d._lock:
        no = d.next_code("capex", "번호", "CX")
        d.upsert("capex", [{"번호": no, "태그번호": tag, "일자": when, "내용": content, "금액": to_num(amount),
                            "연장연수": int(to_num(extend_years, 0) or 0), "등록자": actor, "등록일시": d.now_text()}])
    d.add_history(tag, "자본적 지출", f"{content} · {int(to_num(amount)):,}원" +
                  (f" · 내용연수 {int(to_num(extend_years))}년 연장" if to_num(extend_years) else ""), actor)
    return no


def delete_capex(no: str, actor: str):
    c = d.lookup("capex", no)
    if not c:
        return
    d.delete_where("capex", '"번호"=?', (no,))
    d.add_history(c["태그번호"], "자본적 지출 취소", f"{c['내용']} · {int(c['금액'] or 0):,}원", actor)


# =========================================================
# 회계 분개 (ERP 업로드용)
# =========================================================
ASSET_ACCT = {"비품": ("212", "비품"), "공기구비품": ("210", "공구와기구"), "차량운반구": ("208", "차량운반구"),
              "기계장치": ("206", "기계장치"), "소프트웨어": ("232", "소프트웨어")}
ACCUM_ACCT = {"비품": ("213", "감가상각누계액(비품)"), "공기구비품": ("211", "감가상각누계액(공구와기구)"),
              "차량운반구": ("209", "감가상각누계액(차량운반구)"), "기계장치": ("207", "감가상각누계액(기계장치)")}
DEP_EXP, AMORT_EXP, SUPPLY_EXP = ("818", "감가상각비"), ("840", "무형자산상각비"), ("830", "소모품비")
PAYABLE, CASH, GAIN, LOSS = ("253", "미지급금"), ("103", "보통예금"), ("914", "유형자산처분이익"), ("970", "유형자산처분손실")
JOURNAL_COLS = ["전표일자", "전표번호", "구분", "차대", "계정코드", "계정과목", "금액", "현장", "적요"]


def journal_entries(year: int, month: int) -> list[dict]:
    """해당 월의 취득·자본적지출·감가상각·처분 분개(현장·계정별 합계). 차변 합계 = 대변 합계."""
    first = date(year, month, 1)
    last = d._shift_months_exact(first, 1) - timedelta(days=1)
    rows, seq = [], [0]

    def voucher(kind, site, lines, memo):
        lines = [(dc, acct, round(amt)) for dc, acct, amt in lines if round(amt) != 0]
        if not lines:
            return
        diff = sum(a for dc, _, a in lines if dc == "차변") - sum(a for dc, _, a in lines if dc == "대변")
        if diff:  # 반올림 차이는 마지막 줄에서 맞춤
            dc, acct, amt = lines[-1]
            lines[-1] = (dc, acct, amt + (diff if dc == "대변" else -diff))
        seq[0] += 1
        no = f"{year}{month:02d}-{seq[0]:03d}"
        for dc, (code, name), amt in lines:
            rows.append({"전표일자": last, "전표번호": no, "구분": kind, "차대": dc, "계정코드": code, "계정과목": name,
                         "금액": amt, "현장": site, "적요": memo})

    # 1) 취득 · 자본적 지출
    acq: dict[tuple, list] = {}
    for a in d.load("assets"):
        if a["취득일자"] and first <= a["취득일자"] <= last:
            acq.setdefault((a["현장"], a["계정과목"]), []).append(a["취득가액"] or 0)
    for c in d.load("capex"):
        a = d.get_asset(c["태그번호"])
        if a and c["일자"] and first <= c["일자"] <= last:
            acq.setdefault((a["현장"], a["계정과목"]), []).append(c["금액"] or 0)
    for (site, acc), amts in sorted(acq.items()):
        total = sum(amts)
        debit = SUPPLY_EXP if acc == "소모품" else ASSET_ACCT[acc]
        voucher("취득", site, [("차변", debit, total), ("대변", PAYABLE, total)], f"{acc} {len(amts)}건 취득")

    # 2) 감가상각 (당월분 = 당월까지 누계 - 전월까지 누계)
    cur = {r["태그번호"]: r for r in d.depreciation_schedule(year, month)}
    prev = {r["태그번호"]: r for r in d.depreciation_schedule(year, month - 1)} if month > 1 else {}
    dep: dict[tuple, float] = {}
    for tag, r in cur.items():
        m = r["당기 상각비"] - (prev.get(tag, {}).get("당기 상각비", 0.0))
        if m > 0.5:
            dep[(r["현장"], r["계정과목"])] = dep.get((r["현장"], r["계정과목"]), 0.0) + m
    for (site, acc), amt in sorted(dep.items()):
        if acc == "소프트웨어":
            voucher("감가상각", site, [("차변", AMORT_EXP, amt), ("대변", ASSET_ACCT[acc], amt)], f"{month}월 {acc} 상각")
        else:
            voucher("감가상각", site, [("차변", DEP_EXP, amt), ("대변", ACCUM_ACCT[acc], amt)], f"{month}월 {acc} 감가상각")

    # 3) 처분 · 망실
    disp: dict[tuple, dict] = {}
    for a in d.load("assets"):
        if a["상태"] not in d.CLOSED_STATUSES or not a["처분일자"] or not (first <= a["처분일자"] <= last):
            continue
        if a["계정과목"] == "소모품":
            continue
        dep_ = d.depreciation(a, a["처분일자"])
        g = disp.setdefault((a["현장"], a["계정과목"], a["처분방법"] or a["상태"]),
                            {"n": 0, "cost": 0.0, "accum": 0.0, "cash": 0.0, "gl": 0.0})
        g["n"] += 1
        g["cost"] += dep_["cost"]
        g["accum"] += dep_["accum"]
        g["cash"] += a["처분금액"] or 0
        g["gl"] += (a["처분금액"] or 0) - dep_["nbv"]
    for (site, acc, method), g in sorted(disp.items()):
        lines = []
        if acc == "소프트웨어":
            credit_asset = g["cost"] - g["accum"]  # 무형자산은 상각액을 직접 차감
        else:
            credit_asset = g["cost"]
            lines.append(("차변", ACCUM_ACCT[acc], g["accum"]))
        if g["cash"]:
            lines.append(("차변", CASH, g["cash"]))
        if g["gl"] < 0:
            lines.append(("차변", LOSS, -g["gl"]))
        lines.append(("대변", ASSET_ACCT[acc], credit_asset))
        if g["gl"] > 0:
            lines.append(("대변", GAIN, g["gl"]))
        voucher("처분", site, lines, f"{acc} {g['n']}건 {method}")
    return rows


# =========================================================
# 분석
# =========================================================
def monthly_flow(years: int = 5) -> list[dict]:
    """월별 취득·처분 건수와 금액."""
    start = date(d.today().year - years + 1, 1, 1)
    agg: dict[str, dict] = {}
    for a in d.load("assets"):
        if a["취득일자"] and a["취득일자"] >= start:
            g = agg.setdefault(a["취득일자"].strftime("%Y-%m"), {"취득 건수": 0, "취득 금액": 0.0, "처분 건수": 0, "처분 금액": 0.0})
            g["취득 건수"] += 1
            g["취득 금액"] += a["취득가액"] or 0
        if a["처분일자"] and a["처분일자"] >= start:
            g = agg.setdefault(a["처분일자"].strftime("%Y-%m"), {"취득 건수": 0, "취득 금액": 0.0, "처분 건수": 0, "처분 금액": 0.0})
            g["처분 건수"] += 1
            g["처분 금액"] += a["처분금액"] or 0
    return [{"월": k, **v} for k, v in sorted(agg.items())]


def site_health() -> list[dict]:
    """현장별 유휴율·내용연수 만료율·평균 사용 경과율·누적 수리비."""
    rep = d.repair_costs()
    out = []
    for site in d.site_names():
        items = [a for a in d.active_assets() if a["현장"] == site]
        dep = [a for a in items if d.account_info(a)["cls"] != "소모품"]
        used = [min(1.0, d.months_used(a["취득일자"], d.today()) / (d.asset_life(a) * 12)) for a in dep if d.asset_life(a)]
        out.append({"현장": site, "보유": len(items),
                    "유휴율(%)": round(sum(a["상태"] == "유휴" for a in items) / len(items) * 100, 1) if items else 0,
                    "만료율(%)": round(sum((d.life_remaining_months(a) or 1) <= 0 for a in dep) / len(dep) * 100, 1) if dep else 0,
                    "평균 경과율(%)": round(sum(used) / len(used) * 100, 1) if used else 0,
                    "누적 수리비": round(sum(rep.get(a["태그번호"], 0) for a in items))})
    return out


def age_distribution() -> list[dict]:
    """내용연수 경과율 구간별 자산 수(소모품 제외)."""
    buckets = [("0~25%", 0.25), ("25~50%", 0.5), ("50~75%", 0.75), ("75~100%", 1.0)]
    counts = {name: 0 for name, _ in buckets}
    counts["만료"] = 0
    for a in d.active_assets():
        life = d.asset_life(a)
        if d.account_info(a)["cls"] == "소모품" or not life:
            continue
        r = d.months_used(a["취득일자"], d.today()) / (life * 12)
        if r > 1:
            counts["만료"] += 1
            continue
        for name, top in buckets:
            if r <= top:
                counts[name] += 1
                break
    return [{"경과율": k, "자산 수": v} for k, v in counts.items()]


def disposal_age_by_account() -> list[dict]:
    """계정과목별 평균 실제 사용기간(처분 자산 기준)과 내용연수 비교."""
    agg: dict[str, list] = {}
    for a in d.closed_assets():
        if a["취득일자"] and a["처분일자"]:
            agg.setdefault(a["계정과목"], []).append(d.months_used(a["취득일자"], a["처분일자"]) / 12)
    return [{"계정과목": k, "처분 건수": len(v), "평균 사용(년)": round(sum(v) / len(v), 1),
             "기본 내용연수(년)": d.ACCOUNTS[k]["life"] or None} for k, v in sorted(agg.items())]


def top_repairs(limit: int = 15) -> list[dict]:
    rep = d.repair_costs()
    out = []
    for tag, cost in sorted(rep.items(), key=lambda kv: -kv[1])[:limit]:
        a = d.get_asset(tag) or {}
        out.append({"태그번호": tag, "품명": a.get("품명", ""), "현장": a.get("현장", ""), "누적 수리비": round(cost),
                    "장부가": round(d.book_value(a)) if a else 0, "상태": a.get("상태", "")})
    return out


# =========================================================
# 알림
# =========================================================
def alerts_for(user: dict, pending: dict | None = None, detail: bool = True, rentals: list[dict] | None = None) -> list[dict]:
    """지금 사용자에게 보여 줄 알림. 심각도: crit(지연) / warn(임박·대기) / info.

    각 알림의 '대상'은 바로 처리할 곳입니다: ("req", 종류, 요청번호) · ("rental", 대여번호) · ("plan", 계획번호)
    · ("asset", 태그번호) · ("return", 회수번호) · ("page", 메뉴키).
    detail=False면 결재 대기를 종류별 건수 한 줄로 묶습니다(메일 요약용).
    """
    role, site = user.get("역할"), user.get("현장") or None
    scope = site if role in SITE_ROLES else None
    out = []

    def add(level, cat, title, detail="", menu="", target=None):
        out.append({"심각도": level, "분류": cat, "제목": title, "내용": detail, "메뉴": menu,
                    "대상": target or ("page", menu)})

    for kind, rows in (pending or {}).items():
        if not rows:
            continue
        name = d.KIND_NAMES.get(kind, kind)
        if not detail:
            add("warn", "결재", f"{name} {len(rows)}건이 처리를 기다립니다", menu="inbox")
            continue
        for r in rows:
            if kind == "inc" and r["상태"] == "경위요청":
                add("crit", "경위서", f"경위서 작성 요청: {r.get('품명', '')}", f"{r['요청번호']} · {r.get('태그번호', '')} · {r.get('현장', '')}",
                    "inbox", ("req", kind, r["요청번호"]))
                continue
            title = r.get("품명", "") or r.get("태그번호", "")
            add("warn", f"결재 · {name}", f"{name} {r['상태']}: {str(title)[:40]}",
                f"{r['요청번호']} · {r.get('현장') or r.get('보내는 현장', '')} · {r.get('요청일시', '')}",
                "inbox", ("req", kind, r["요청번호"]))
    for r in rentals or []:
        add("warn", "대여 신청", f"대여 {'승인' if r['상태'] == '신청' else '인계'} 대기: {r['품명']} ({r['신청자명']})",
            f"{r['대여번호']} · {r['시작일']} ~ {r['반납예정일']}", "rental", ("rental", r["대여번호"]))
    if role in (ADMIN,) + SITE_ROLES:
        for r in overdue_rentals(scope):
            add("crit", "대여", f"반납 지연: {r['품명']} ({r['신청자명']})", f"반납 예정 {r['반납예정일']} · {r['대여번호']}", "rental",
                ("rental", r["대여번호"]))
        for p in plans_due(site=scope):
            add("crit" if p["상태"] == "지연" else "warn", "점검",
                f"{p['구분']} {'지연' if p['상태'] == '지연' else '임박'}: {p['품명']} · {p['점검명']}",
                f"{p['태그번호']} · 예정일 {p['다음점검일']} ({p['남은 일수']:+d}일)", "maint", ("plan", p["계획번호"]))
        for a in warranty_expiring(site=scope)[:30]:
            add("info", "보증", f"보증 만료 임박: {a['품명']}", f"{a['태그번호']} · {a['보증만료일']}", "maint", ("asset", a["태그번호"]))
        for r in get_returns(scope):
            if r["상태"] == "진행중":
                left = len([i for i in return_items(r["회수번호"]) if i["결과"] == "미처리"])
                add("warn", "퇴사 회수", f"{r['이름']} 퇴사 회수 미처리 {left}건", f"{r['회수번호']} · 퇴사일 {r['퇴사일']}", "people",
                    ("return", r["회수번호"]))
    if role in (ADMIN, FIN):
        for x in licenses_expiring():
            add("crit" if x["만료까지(일)"] < 0 else "warn", "라이선스", f"갱신 필요: {x['소프트웨어명']}",
                f"만료 {x['만료일']} · 갱신비 {int(x['갱신비용'] or 0):,}원", "license")
        for x in licenses_expiring(days=10**6):
            if x["남음"] < 0:
                add("crit", "라이선스", f"초과 사용: {x['소프트웨어명']}", f"보유 {x['보유수량']} · 사용 {x['사용']}", "license")
        for x in leases_expiring(site=scope):
            add("warn", "렌탈·리스", f"계약 만료 임박: {x['품명']} ({x['현장']})", f"{x['업체']} · 종료 {x['종료일']}", "lease")
        for b in budget_status(d.today().year):
            if b["예산"] and b["잔액"] < 0:
                add("crit", "예산", f"{b['현장']} 예산 초과", f"예산 {int(b['예산']):,}원 · 초과 {int(-b['잔액']):,}원", "budget")
            elif b["예산"] and b["집행률(%)"] and b["집행률(%)"] >= 90:
                add("warn", "예산", f"{b['현장']} 예산 {b['집행률(%)']}% 집행", "", "budget")
    my_rentals = [r for r in d.load("rentals") if r["신청자"] == user.get("아이디") and r["상태"] == "대여중"]
    for r in my_rentals:
        days = (r["반납예정일"] - d.today()).days if r["반납예정일"] else 0
        if days <= 1:
            if any(x["대상"] == ("rental", r["대여번호"]) for x in out):
                continue  # 이미 '반납 지연'으로 들어간 건
            add("crit" if days < 0 else "warn", "대여", f"내 대여 반납 {'지연' if days < 0 else '예정'}: {r['품명']}",
                f"반납 예정 {r['반납예정일']} · {r['대여번호']}", "rental", ("rental", r["대여번호"]))
    order = {"crit": 0, "warn": 1, "info": 2}
    return sorted(out, key=lambda x: order[x["심각도"]])


def digest_text(alerts: list[dict], title: str = "자산관리 AI 알림") -> str:
    if not alerts:
        return f"[{title}] {d.today()} 처리할 알림이 없습니다."
    mark = {"crit": "🔴", "warn": "🟡", "info": "🔵"}
    lines = [f"[{title}] {d.today()} · {len(alerts)}건"]
    for a in alerts[:60]:
        lines.append(f"{mark[a['심각도']]} [{a['분류']}] {a['제목']}" + (f" — {a['내용']}" if a["내용"] else ""))
    if len(alerts) > 60:
        lines.append(f"… 외 {len(alerts) - 60}건")
    return "\n".join(lines)


def send_email(subject: str, body: str, to: list[str] | None = None) -> str:
    s = d.get_settings()
    server, account = clean_str(s.get("메일 서버")), clean_str(s.get("메일 계정"))
    to = to or [x.strip() for x in clean_str(s.get("알림 받는 메일")).split(",") if x.strip()]
    if not server or not account or not to:
        raise DataError("메일 서버·계정·받는 메일을 설정에서 입력하세요.")
    msg = MIMEText(body, "plain", "utf-8")
    msg["Subject"], msg["From"], msg["To"] = subject, account, ", ".join(to)
    try:
        with smtplib.SMTP(server, int(to_num(s.get("메일 포트"), 587) or 587), timeout=20) as smtp:
            smtp.starttls()
            smtp.login(account, clean_str(s.get("메일 비밀번호")))
            smtp.sendmail(account, to, msg.as_string())
    except Exception as exc:
        raise DataError(f"메일을 보내지 못했습니다: {exc}") from exc
    return f"메일 {len(to)}명"


def send_webhook(text: str, url: str | None = None) -> str:
    url = clean_str(url or d.get_settings().get("웹훅 주소"))
    if not url:
        raise DataError("웹훅 주소를 설정에서 입력하세요.")
    if requests is None:
        raise DataError("requests 패키지가 없습니다.")
    payload = {"content": text} if "discord" in url else {"text": text}
    try:
        r = requests.post(url, data=json.dumps(payload), headers={"Content-Type": "application/json"}, timeout=15)
        r.raise_for_status()
    except Exception as exc:
        raise DataError(f"웹훅을 보내지 못했습니다: {exc}") from exc
    return "웹훅"


def send_digest(alerts: list[dict], channels: tuple = ("email", "webhook")) -> list[str]:
    """알림 요약을 메일·웹훅으로 보냅니다. 설정되지 않은 채널은 건너뜁니다."""
    text = digest_text(alerts)
    done, s = [], d.get_settings()
    for ch in channels:
        try:
            if ch == "email" and clean_str(s.get("메일 서버")):
                done.append(send_email(f"[자산관리 AI] {d.today()} 알림 {len(alerts)}건", text))
            elif ch == "webhook" and clean_str(s.get("웹훅 주소")):
                done.append(send_webhook(text))
            else:
                continue
            result = "성공"
        except DataError as exc:
            result = str(exc)
            done.append(f"{ch} 실패: {exc}")
        d.insert_bulk("notify_log", [{"일시": d.now_text(), "채널": ch, "수신": "", "건수": len(alerts), "결과": result}])
    return done


# =========================================================
# 법인차량 관리 프로그램 연동 (렌트 → 리스, 정비 → 점검 기록, 운행일지 → 월별 KM)
# =========================================================
def sync_vehicle_extras(vehicles: list[dict], actor: str) -> dict:
    res = {"leases": 0, "maint": 0}
    linked = {a["연결차량번호"]: a for a in d.load("assets") if a["연결차량번호"]}
    lease_by_car = {x["연결차량번호"]: x for x in d.load("leases") if x["연결차량번호"]}
    sites = set(d.site_names())
    for v in vehicles:
        no = clean_str(v.get("차량번호"))
        kind = clean_str(v.get("구분"))
        if kind == "렌트" and no not in lease_by_car and clean_str(v.get("현장")) in sites:
            save_lease({"구분": "렌탈", "품명": f"{clean_str(v.get('차종'))} (업무용 렌트)", "규격": f"차량번호 {no}",
                        "현장": clean_str(v.get("현장")), "업체": "렌터카", "시작일": to_date(v.get("차량 구매일")),
                        "월요금": to_num(v.get("월 렌트료"), 0) or 0, "담당": clean_str(v.get("관리자")),
                        "연결차량번호": no, "비고": "법인차량 연동"}, actor)
            res["leases"] += 1
        a = linked.get(no)
        last = to_date(v.get("최근정비일"))
        if a and last and not any(r["태그번호"] == a["태그번호"] and r["일자"] == last for r in d.load("maint_logs")):
            d.add_maint_log({"태그번호": a["태그번호"], "일자": last, "구분": "정기점검", "업체": "",
                             "내용": f"차량 정비 (법인차량 연동 · 주행 {int(to_num(v.get('누적주행거리'), 0) or 0):,}km)",
                             "결과": "정상"}, actor)
            res["maint"] += 1
    return res


def import_vehicle_km(source, year: int | None = None) -> int:
    """운행일지.xlsx의 '월별요약' 시트(차량번호, 01월 KM ~ 12월 KM)를 읽어 차량별 월 주행거리를 저장합니다."""
    from openpyxl import load_workbook
    if isinstance(source, (str, Path)):
        path = Path(clean_str(source).strip('"'))
        if not path.exists():
            raise DataError(f"운행일지 파일을 찾을 수 없습니다: {path}")
        source = path
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        if "월별요약" not in wb.sheetnames:
            raise DataError("운행일지 파일에 '월별요약' 시트가 없습니다.")
        rows = wb["월별요약"].iter_rows(values_only=True)
        header = [clean_str(h) for h in (next(rows, None) or [])]
        if "차량번호" not in header:
            raise DataError("'월별요약' 시트에 '차량번호' 열이 없습니다.")
        year = year or d.today().year
        out = []
        for r in rows:
            rec = dict(zip(header, r))
            car = clean_str(rec.get("차량번호"))
            if not car:
                continue
            for m in range(1, 13):
                km = to_num(rec.get(f"{m:02d}월 KM"), 0) or 0
                out.append({"차량번호": car, "연월": f"{year}-{m:02d}", "km": km})
    finally:
        wb.close()
    d.upsert("vehicle_km", out)
    return len({r["차량번호"] for r in out})


def vehicle_km(car_no: str) -> list[dict]:
    return sorted([r for r in d.load("vehicle_km") if r["차량번호"] == car_no], key=lambda r: r["연월"])


# =========================================================
# 예시 데이터 (data.create_sample_data에서 호출)
# =========================================================
_SURNAMES = list("김이박최정강조윤장임한오서신권황안송류전홍고문양손배백허유남심노하곽성차주우구민진나")
_GIVEN = ["민수", "지훈", "서연", "우성", "지민", "세훈", "나래", "태호", "지수", "도현", "재현", "유진", "성민", "수아",
          "준호", "가영", "동욱", "하늘", "보람", "지아", "현우", "은지", "상훈", "미영", "정우", "수빈", "영호", "하은",
          "건우", "예린", "승민", "다은", "태윤", "소영", "민재", "윤서", "진혁", "채원", "경민", "혜진"]
SAMPLE_DEPTS_BY_KIND = {"본사": ["경영지원팀", "재무팀", "인사팀", "구매팀", "설계팀", "IT팀"],
                        "현장": ["공무팀", "공사팀", "설비팀", "품질팀", "안전팀", "운영팀"]}


def sample_employees(sites: list[str], rnd, total: int) -> dict[str, list[dict]]:
    """현장별 예시 직원(사번·이름·부서)을 만들어 저장하고 현장별 목록을 돌려줍니다."""
    names, pool, recs = set(), {s: [] for s in sites}, []
    n = 0
    while len(recs) < total:
        name = rnd.choice(_SURNAMES) + rnd.choice(_GIVEN)
        if name in names and rnd.random() < 0.9:
            continue
        names.add(name)
        site = sites[len(recs) % len(sites)]
        n += 1
        depts = SAMPLE_DEPTS_BY_KIND["본사" if site == sites[0] else "현장"]
        e = {"사번": f"E{n:05d}", "이름": name, "현장": site, "부서": rnd.choice(depts),
             "직급": rnd.choice(["사원", "주임", "대리", "과장", "차장", "부장"]), "이메일": f"user{n:05d}@example.com",
             "연락처": f"010-{rnd.randint(1000, 9999)}-{rnd.randint(1000, 9999)}", "상태": "재직",
             "입사일": date(rnd.randint(2008, 2025), rnd.randint(1, 12), rnd.randint(1, 28)), "퇴사일": None, "비고": d.SAMPLE_MARK}
        recs.append(e)
        pool[site].append(e)
    d.insert_bulk("employees", recs)
    return pool


def sample_extras(active: list[dict], sites: list[str], rnd, actor: str, pool: dict[str, list[dict]]):
    """새 기능용 예시 데이터: 업체·예산·점검·라이선스·렌탈리스·대여·자본적지출·퇴사 회수·경위서 요청."""
    t = d.today()
    mark = d.SAMPLE_MARK
    # 1) 공급업체
    vendors = ["OO상사", "한빛정보", "대성산업", "미래테크", "에이스장비", "에어로텍", "스카이맵", "드론원", "가나정보통신", "세진중기"]
    d.insert_bulk("vendors", [{"업체명": v, "사업자번호": f"{rnd.randint(100, 999)}-{rnd.randint(10, 99)}-{rnd.randint(10000, 99999)}",
                               "담당자": rnd.choice(_SURNAMES) + rnd.choice(_GIVEN), "연락처": f"02-{rnd.randint(100, 999)}-{rnd.randint(1000, 9999)}",
                               "이메일": f"sales{i}@vendor.example.com", "취급품목": rnd.choice(["IT기기", "사무가구", "측량장비", "중장비", "드론"]),
                               "상태": "거래중", "비고": mark, "등록일시": d.now_text()} for i, v in enumerate(vendors)])
    # 2) 예산 (올해·작년, 현장별 취득 실적의 1.1~1.4배)
    by_site_year: dict[tuple, float] = {}
    for a in active:
        y = a["취득일자"].year
        if y >= t.year - 1:
            by_site_year[(y, a["현장"])] = by_site_year.get((y, a["현장"]), 0) + a["취득가액"]
    budgets = []
    for y in (t.year - 1, t.year):
        for s in sites:
            base = by_site_year.get((y, s), 50_000_000) * (12 / max(1, t.month) if y == t.year else 1)
            budgets.append({"연도": y, "현장": s, "예산액": round(base * rnd.uniform(1.1, 1.4), -7), "비고": mark})
    d.insert_bulk("budgets", budgets)
    # 3) 공용 장비·보증·세무 내용연수는 자산 생성 때 반영됨. 점검 계획·기록
    plan_rules = {"전동 지게차": ("지게차 안전검사", "법정검사", 12), "공기압축기": ("압력용기 정기검사", "법정검사", 12),
                  "발전기": ("발전기 정기점검", "정기점검", 6), "컨베이어 라인": ("컨베이어 정기점검", "정기점검", 6),
                  "현장 승합차": ("차량 정기점검", "정기점검", 6), "화물차": ("차량 정기점검", "정기점검", 6),
                  "토탈스테이션": ("측량기 검교정", "검교정", 12), "레이저 레벨기": ("측량기 검교정", "검교정", 12)}
    cost_by_kind = {"법정검사": (150_000, 400_000), "정기점검": (80_000, 350_000), "검교정": (200_000, 450_000)}
    plans, logs, pn, ln = [], [], 0, 0
    for a in active:
        rule = plan_rules.get(a["품명"])
        if not rule:
            continue
        name, kind, cyc = rule
        pn += 1
        last = t - timedelta(days=rnd.randint(10, int(cyc * 30.5 * 1.12)))
        last = max(last, a["취득일자"])
        nxt = d._shift_months_exact(last, cyc)
        no = f"MP-S{pn:05d}"
        plans.append({"계획번호": no, "태그번호": a["태그번호"], "점검명": name, "구분": kind, "주기개월": cyc,
                      "최근점검일": last, "다음점검일": nxt, "담당업체": rnd.choice(vendors), "활성": "Y", "비고": mark})
        when = last
        for _ in range(rnd.randint(1, 3)):
            if when < a["취득일자"]:
                break
            ln += 1
            lo, hi = cost_by_kind[kind]
            logs.append({"기록번호": f"ML-S{ln:06d}", "태그번호": a["태그번호"], "일자": when, "구분": kind, "내용": name,
                         "업체": rnd.choice(vendors), "비용": round(rnd.uniform(lo, hi), -4),
                         "결과": "정상" if rnd.random() > 0.08 else "조치필요", "계획번호": no, "등록자": actor, "등록일시": d.now_text()})
            when = d._shift_months_exact(when, -cyc)
    for a in active:  # 수리 기록 (약 4%), 일부는 누적 수리비가 장부가를 넘김
        if a["계정과목"] in ("기계장치", "공기구비품", "비품") and rnd.random() < 0.04:
            for _ in range(rnd.randint(1, 3)):
                ln += 1
                day = a["취득일자"] + timedelta(days=rnd.randint(30, max(31, (t - a["취득일자"]).days)))
                logs.append({"기록번호": f"ML-S{ln:06d}", "태그번호": a["태그번호"], "일자": min(day, t), "구분": "수리",
                             "내용": rnd.choice(["부품 교체", "모터 수리", "메인보드 교체", "유압 계통 수리", "배터리 교체"]),
                             "업체": rnd.choice(vendors), "비용": round(a["취득가액"] * rnd.uniform(0.05, 0.45), -4),
                             "결과": "정상", "계획번호": "", "등록자": actor, "등록일시": d.now_text()})
    d.insert_bulk("maint_plans", plans)
    d.insert_bulk("maint_logs", logs)
    # 4) 소프트웨어 라이선스
    emps = [e for v in pool.values() for e in v]
    pcs = {}
    for a in active:
        if a["품명"] in ("노트북", "데스크탑 PC", "워크스테이션") and a["사용자사번"]:
            pcs.setdefault(a["사용자사번"], a["태그번호"])
    lic_defs = [("Microsoft 365 E3", "Enterprise", "구독", 300, t + timedelta(days=40), 9_800_000, 0.93),
                ("AutoCAD", "LT 2026", "구독", 25, t + timedelta(days=200), 12_500_000, 1.0),
                ("한컴오피스 2024", "기업용", "영구", 150, None, 0, 0.7),
                ("CATIA V5", "MD2", "영구", 6, None, 0, 0.83),
                ("V3 Endpoint Security", "9.0", "구독", 400, t + timedelta(days=15), 6_000_000, 0.88),
                ("Adobe Acrobat Pro", "팀용", "구독", 40, t - timedelta(days=5), 3_600_000, 1.05)]
    lics, assigns, an = [], [], 0
    for i, (name, ed, typ, qty, exp, cost, fill) in enumerate(lic_defs, start=1):
        no = f"LC-S{i:03d}"
        lics.append({"라이선스번호": no, "소프트웨어명": name, "에디션": ed, "유형": typ, "보유수량": qty,
                     "계약일": (exp or t) - timedelta(days=365), "만료일": exp, "갱신비용": cost, "공급업체": "가나정보통신",
                     "연결태그": "", "현장": "", "상태": "사용", "비고": mark})
        for e in rnd.sample(emps, min(len(emps), int(qty * fill))):
            an += 1
            assigns.append({"할당번호": f"LA-S{an:05d}", "라이선스번호": no, "사용자": e["이름"], "사번": e["사번"], "부서": e["부서"],
                            "현장": e["현장"], "PC태그": pcs.get(e["사번"], ""), "할당일": t - timedelta(days=rnd.randint(10, 700)),
                            "해제일": None, "상태": "사용중"})
    d.insert_bulk("licenses", lics)
    d.insert_bulk("license_assign", assigns)
    # 5) 렌탈·리스
    leases, n = [], 0
    for s in sites:
        for item, typ, (lo, hi), corp in [("컬러 복합기", "렌탈", (150_000, 260_000), "신도리코"), ("정수기", "렌탈", (28_000, 45_000), "코웨이"),
                                          ("업무용 차량 장기렌트", "렌탈", (600_000, 900_000), "SK렌터카"),
                                          ("업무용 차량 장기렌트", "렌탈", (600_000, 900_000), "롯데렌탈")]:
            n += 1
            start = t - timedelta(days=rnd.randint(100, 1000))
            leases.append({"계약번호": f"LS-S{n:03d}", "구분": typ, "품명": item, "규격": "", "현장": s, "업체": corp,
                           "시작일": start, "종료일": start + timedelta(days=365 * rnd.choice([2, 3, 4])), "월요금": round(rnd.uniform(lo, hi), -3),
                           "보증금": 0, "자동연장": "Y" if typ == "렌탈" and rnd.random() < 0.5 else "", "담당": "",
                           "상태": "사용중", "연결차량번호": "", "비고": mark})
    leases.append({"계약번호": "LS-S900", "구분": "리스", "품명": "서버 운용리스", "규격": "Dell PowerEdge R760 ×4", "현장": sites[0],
                   "업체": "KB캐피탈", "시작일": t - timedelta(days=1060), "종료일": t + timedelta(days=20), "월요금": 1_250_000,
                   "보증금": 5_000_000, "자동연장": "", "담당": "", "상태": "사용중", "연결차량번호": "", "비고": mark})
    for x in leases:
        if x["종료일"] < t:
            x["상태"] = "종료"
    d.insert_bulk("leases", leases)
    # 6) 공용 장비 대여
    shared = [a for a in active if a["공용"] == "Y" and a["상태"] in ("사용중", "유휴")]
    rentals, rn = [], 0

    def rental(a, start, days, status, who):
        nonlocal rn
        rn += 1
        end = start + timedelta(days=days)
        r = {"대여번호": f"RT-S{rn:05d}", "상태": status, "태그번호": a["태그번호"], "품명": a["품명"], "현장": a["현장"],
             "신청자": who.get("아이디", who.get("사번", "")), "신청자명": who["이름"], "부서": who.get("부서", ""),
             "시작일": start, "반납예정일": end, "사용목적": rnd.choice(["현장 측량", "고객사 미팅", "외부 교육", "공정 점검", "출장"]),
             "신청일시": f"{start - timedelta(days=3)} 09:00", "승인자": "", "승인일시": "", "대여일시": "", "반납일시": "",
             "반납상태": "", "메모": "", "처리이력": f"[{start - timedelta(days=3)} 09:00] 대여 신청 · {mark} {who['이름']}"}
        if status in ("승인", "대여중", "반납완료"):
            r["승인자"], r["승인일시"] = f"{mark} 현장 담당", f"{start - timedelta(days=2)} 10:00"
        if status in ("대여중", "반납완료"):
            r["대여일시"] = f"{start} 09:00"
        if status == "반납완료":
            r["반납일시"], r["반납상태"] = f"{end} 17:00", "정상" if rnd.random() > 0.05 else "이상"
        rentals.append(r)
    used = set()
    for a in rnd.sample(shared, min(len(shared), 80)):
        e = rnd.choice(pool[a["현장"]]) if pool.get(a["현장"]) else {"이름": "예시 직원"}
        st = rnd.choice(["반납완료"] * 6 + ["대여중", "승인", "신청"])
        if st == "반납완료":
            rental(a, t - timedelta(days=rnd.randint(15, 180)), rnd.randint(1, 5), st, e)
        elif st == "대여중":
            rental(a, t - timedelta(days=rnd.randint(1, 6)), rnd.randint(2, 7), st, e)
        else:
            rental(a, t + timedelta(days=rnd.randint(2, 20)), rnd.randint(1, 4), st, e)
        used.add(a["태그번호"])
    site1 = sites[1] if len(sites) > 1 else sites[0]
    staff = {"아이디": "staff", "이름": "한직원", "부서": "설계팀"}
    mine = [a for a in shared if a["현장"] == site1 and a["태그번호"] not in used][:3]
    if mine:  # 체험 계정(staff)의 대여: 지연 1건, 신청 1건
        rental(mine[0], t - timedelta(days=6), 3, "대여중", staff)
        if len(mine) > 1:
            rental(mine[1], t + timedelta(days=3), 2, "신청", staff)
    d.insert_bulk("rentals", rentals)
    # 7) 자본적 지출 (기계장치 일부)
    capex, cn = [], 0
    for a in [x for x in active if x["계정과목"] == "기계장치" and (t - x["취득일자"]).days > 400][:14]:
        cn += 1
        capex.append({"번호": f"CX-S{cn:03d}", "태그번호": a["태그번호"], "일자": a["취득일자"] + timedelta(days=rnd.randint(300, (t - a["취득일자"]).days)),
                      "내용": rnd.choice(["엔진 오버홀", "용량 증설", "제어반 교체", "구동부 개량"]),
                      "금액": round(a["취득가액"] * rnd.uniform(0.08, 0.25), -4), "연장연수": rnd.choice([0, 1, 2]),
                      "등록자": actor, "등록일시": d.now_text()})
    d.insert_bulk("capex", capex)
    # 8) 퇴사 회수 진행 중 1건 + 경위서 요청 2건 (체험 계정 현장)
    holders = {}
    for a in active:
        if a["현장"] == site1 and a["사용자사번"]:
            holders.setdefault(a["사용자사번"], []).append(a)
    emp_no = next((k for k, v in holders.items() if len(v) >= 3), None)
    if emp_no:
        rc = start_return(emp_no, t + timedelta(days=7), actor)
        first = return_items(rc)[0]
        process_return_item(rc, first["태그번호"], "회수", "", "반납 확인", actor)
    for a in [x for x in active if x["현장"] == site1 and x["상태"] == "사용중" and x["계정과목"] in ("비품", "공기구비품")][5:7]:
        d.request_statement(a["태그번호"], "재물조사 미발견 (예시)", actor)
