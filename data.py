"""
자산관리 AI - 데이터 계층

저장소: data/자산관리.db (SQLite)
- 보유 자산 수만 개 + 처분 자산 수십만 개를 다루기 위해 엑셀 대신 SQLite를 씁니다.
  (엑셀 20만 행은 저장 한 번에 1분가량 걸립니다.)
- 모든 표는 설정 화면에서 엑셀(xlsx)·CSV로 내보낼 수 있고, 예전 버전의 data/*.xlsx가 있으면
  처음 실행할 때 자동으로 옮겨 옵니다.
- 자주 쓰는 표(자산, 현장, 요청 등)는 메모리에 캐시하고, 저장할 때 바뀐 행만 DB와 캐시에 함께 반영합니다.
- 모든 쓰기는 하나의 잠금(RLock) 안에서 처리되어 여러 브라우저에서 동시에 작업해도 안전합니다.

화면 코드(admin.py, user.py, workflow.py)는 이 모듈의 함수만 호출합니다.
"""
from __future__ import annotations

import csv
import io
import random
import re
import sqlite3
import threading
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# =========================================================
# 경로 / 기준 정보
# =========================================================
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DB_FILE = DATA_DIR / "자산관리.db"

# 법인차량 관리 프로그램과 같은 현장 구성으로 시작합니다.
DEFAULT_SITES = [
    {"현장명": "본사", "구분": "본사", "표시순서": 1},
    {"현장명": "평택", "구분": "현장", "표시순서": 2},
    {"현장명": "동탄", "구분": "현장", "표시순서": 3},
    {"현장명": "탕정", "구분": "현장", "표시순서": 4},
]

# 계정과목: 이름 → 코드/자산구분/기본 내용연수/기본 상각방법
ACCOUNTS = {
    "비품": {"code": "FUR", "cls": "유형자산", "life": 5, "method": "정률"},
    "공기구비품": {"code": "TLS", "cls": "유형자산", "life": 5, "method": "정률"},
    "차량운반구": {"code": "VEH", "cls": "유형자산", "life": 5, "method": "정률"},
    "기계장치": {"code": "MCH", "cls": "유형자산", "life": 8, "method": "정률"},
    "소프트웨어": {"code": "SW", "cls": "무형자산", "life": 5, "method": "정액"},
    "소모품": {"code": "CON", "cls": "소모품", "life": 0, "method": "-"},
}
ACCOUNT_NAMES = list(ACCOUNTS.keys())
CLASSES = ["유형자산", "무형자산", "소모품"]

ASSET_STATUSES = ["사용중", "유휴", "수리중", "처분대기", "망실", "처분완료"]
CLOSED_STATUSES = {"망실", "처분완료"}
DISPOSAL_METHODS = ["폐기", "매각", "기증", "사내 경매", "망실"]

FLOWS = {
    "pr": ["요청", "승인", "발주", "검수완료", "등록완료"],
    "move": ["요청", "승인", "완료"],
    "inc": ["접수", "조사완료", "종결"],
    "dsp": ["신청", "승인", "집행완료", "제적완료"],
}
KIND_NAMES = {"pr": "구매요청", "move": "이동신청", "inc": "분실·파손", "dsp": "처분"}
KIND_PREFIX = {"pr": "PR", "move": "MV", "inc": "IN", "dsp": "DS"}
REJECTED = "반려"

STORAGE_DEVICE_PATTERN = re.compile(r"노트북|PC|데스크탑|컴퓨터|워크스테이션|서버|태블릿|NAS|SSD|HDD|저장", re.I)

DEFAULT_SETTINGS = {
    "회사명": ("", "QR 라벨 상단에 표시할 회사 이름"),
    "관리자명": ("자산관리자", "관리자 화면에서 처리자로 기록되는 이름"),
    "관리자 PIN": ("", "비워 두면 관리자 페이지에 PIN 없이 들어갑니다"),
    "자본화 기준금액": ("1000000", "단가가 이 금액 미만이면 소모품(비용)으로 분류를 제안합니다"),
    "QR 기본주소": ("", "예: http://192.168.0.10:8502  비우면 QR에 태그번호만 담깁니다"),
    "차량목록 경로": ("", "법인차량 관리 프로그램의 차량목록.xlsx 전체 경로"),
    "Ollama 주소": ("http://127.0.0.1:11434", "AI 요청 분석에 쓰는 Ollama 서버"),
    "Ollama 모델": ("qwen2.5vl:3b", "AI 요청 분석 모델"),
}

# =========================================================
# 테이블 정의
#   key: 기본키 열, cache: 메모리 캐시 여부, index: 추가 색인
#   legacy: 예전 엑셀 버전의 (파일, 시트) — 자동 이전용
# =========================================================
TABLES = {
    "assets": {
        "label": "자산목록", "key": ["태그번호"], "cache": True, "legacy": ("자산목록.xlsx", "자산목록"),
        "cols": ["태그번호", "품명", "규격", "시리얼번호", "계정과목", "자산구분", "현장", "관리부서", "사용자",
                 "설치위치", "취득일자", "취득가액", "내용연수", "상각방법", "잔존가액", "공급업체", "상태",
                 "이전상태", "수리견적", "처분일자", "처분방법", "처분처", "처분금액", "처분손익", "처분요청번호",
                 "연결차량번호", "구매요청번호", "등록일시", "비고"],
        "dates": ["취득일자", "처분일자"],
        "nums": ["취득가액", "내용연수", "잔존가액", "수리견적", "처분금액", "처분손익"],
    },
    "sites": {
        "label": "현장목록", "key": ["현장명"], "cache": True, "legacy": ("현장목록.xlsx", "현장목록"),
        "cols": ["현장명", "구분", "담당자", "담당자 연락처", "주소", "표시순서", "비고"],
        "dates": [], "nums": ["표시순서"],
    },
    "settings": {
        "label": "설정", "key": ["항목"], "cache": True, "legacy": ("현장목록.xlsx", "설정"),
        "cols": ["항목", "값", "설명"], "dates": [], "nums": [],
    },
    "pr": {
        "label": "구매요청", "key": ["요청번호"], "cache": True, "legacy": ("자산요청.xlsx", "구매요청"),
        "cols": ["요청번호", "상태", "요청일시", "현장", "요청자", "품명", "규격", "계정과목", "수량", "예상단가",
                 "필요일", "사용예정자", "구매사유", "검토의견",
                 "견적1 업체", "견적1 단가", "견적1 조건", "견적2 업체", "견적2 단가", "견적2 조건",
                 "견적3 업체", "견적3 단가", "견적3 조건", "선정업체", "선정단가", "선정사유",
                 "검수 규격", "검수 입고수량", "검수 파손", "검수 시리얼", "검수 메모", "검수자", "검수일시",
                 "등록 태그번호", "처리이력"],
        "dates": ["필요일"],
        "nums": ["수량", "예상단가", "견적1 단가", "견적2 단가", "견적3 단가", "선정단가", "검수 입고수량"],
    },
    "move": {
        "label": "이동신청", "key": ["요청번호"], "cache": True, "legacy": ("자산요청.xlsx", "이동신청"),
        "cols": ["요청번호", "상태", "요청일시", "요청자", "태그번호", "품명", "보내는 현장", "받는 현장",
                 "받는 부서", "새 사용자", "사유", "승인자", "승인일시", "인수자", "인수일시", "처리이력"],
        "dates": [], "nums": [],
    },
    "inc": {
        "label": "분실파손", "key": ["요청번호"], "cache": True, "legacy": ("자산요청.xlsx", "분실파손"),
        "cols": ["요청번호", "상태", "요청일시", "신고자", "태그번호", "품명", "현장", "구분", "발생일", "발생장소",
                 "관련자", "경위", "경찰신고번호", "고의과실", "변상여부", "변상금액", "처리방법", "수리견적",
                 "보험청구", "조사의견", "조사자", "조사일시", "종결일시", "처리이력"],
        "dates": ["발생일"], "nums": ["변상금액", "수리견적"],
    },
    "dsp": {
        "label": "처분", "key": ["요청번호"], "cache": True, "legacy": ("자산요청.xlsx", "처분"),
        "cols": ["요청번호", "상태", "요청일시", "신청자", "현장", "태그번호", "품명", "불용사유", "희망 처분방법",
                 "상세사유", "위원회 의견", "승인일시", "처분방법", "집행일", "처분처", "처분금액", "보안파기 확인",
                 "파기증명번호", "장부가 합계", "처분손익", "제적일시", "처리이력"],
        "dates": ["집행일"], "nums": ["처분금액", "장부가 합계", "처분손익"],
    },
    "history": {
        "label": "자산이력", "key": None, "cache": False, "legacy": ("자산이력.xlsx", "자산이력"),
        "cols": ["일시", "태그번호", "구분", "내용", "처리자"], "dates": [], "nums": [],
        "index": [["태그번호"], ["일시"]],
    },
    "audits": {
        "label": "조사회차", "key": ["회차번호"], "cache": True, "legacy": ("재물조사.xlsx", "조사회차"),
        "cols": ["회차번호", "회차명", "기준일", "대상", "상태", "시작자", "시작일시", "마감일시",
                 "장부수량", "확인", "미발견", "유휴"],
        "dates": ["기준일"], "nums": ["장부수량", "확인", "미발견", "유휴"],
    },
    "audit_items": {
        "label": "조사결과", "key": ["회차번호", "태그번호"], "cache": False, "legacy": ("재물조사.xlsx", "조사결과"),
        "cols": ["회차번호", "태그번호", "품명", "현장", "설치위치", "결과", "유휴", "메모", "확인자", "확인일시"],
        "dates": [], "nums": [], "index": [["회차번호", "현장"]],
    },
    "closings": {
        "label": "결산마감", "key": ["결산월"], "cache": True, "legacy": ("결산마감.xlsx", "마감이력"),
        "cols": ["결산월", "마감일시", "마감자", "자산수", "취득가액", "당기상각비", "상각누계액", "장부가액"],
        "dates": [], "nums": ["자산수", "취득가액", "당기상각비", "상각누계액", "장부가액"],
    },
    "closing_items": {
        "label": "결산명세", "key": ["결산월", "태그번호"], "cache": False, "legacy": None,
        "cols": ["결산월", "태그번호", "품명", "계정과목", "현장", "취득일자", "취득가액", "내용연수", "상각방법",
                 "기초 장부가", "당기 취득", "당기 상각비", "월 상각비", "상각누계액", "기말 장부가", "경과월", "비고"],
        "dates": ["취득일자"],
        "nums": ["취득가액", "내용연수", "기초 장부가", "당기 취득", "당기 상각비", "월 상각비", "상각누계액", "기말 장부가", "경과월"],
    },
}
for _spec in TABLES.values():
    _spec["_text"] = [c for c in _spec["cols"] if c not in _spec["dates"] and c not in _spec["nums"]]
    _spec["_numset"] = set(_spec["nums"])

REQUEST_TABLES = ["pr", "move", "inc", "dsp"]

MONEY_COLS = {"취득가액", "잔존가액", "수리견적", "처분금액", "처분손익", "예상단가", "견적1 단가", "견적2 단가",
              "견적3 단가", "선정단가", "변상금액", "장부가 합계", "당기상각비", "상각누계액", "장부가액",
              "기초 장부가", "당기 취득", "당기 상각비", "월 상각비", "기말 장부가"}


class DataError(Exception):
    """화면에 그대로 보여줄 수 있는 저장/검증 오류."""


_lock = threading.RLock()
_conn: sqlite3.Connection | None = None
_cache: dict[str, dict] = {}
_ver_bump: dict[str, int] = {}
_derived: dict = {}


# =========================================================
# 기본 변환 함수
# =========================================================
def clean_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value != value:
        return ""
    return str(value).strip()


def to_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        if value != value or value <= 0:
            return None
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    text = clean_str(value)[:10].replace(".", "-").replace("/", "-")
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return None


def to_num(value, default=None):
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        if value != value:
            return default
        return int(value) if float(value).is_integer() else float(value)
    digits = re.sub(r"[^0-9.\-]", "", str(value))
    if digits in ("", "-", ".", "-."):
        return default
    try:
        number = float(digits)
    except ValueError:
        return default
    return int(number) if number.is_integer() else number


def today() -> date:
    return date.today()


def now_text() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def _q(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


# =========================================================
# SQLite 저장소
# =========================================================
def _db() -> sqlite3.Connection:
    global _conn
    if _conn is None:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        _conn = sqlite3.connect(DB_FILE, check_same_thread=False, timeout=30)
        _conn.execute("PRAGMA journal_mode=WAL")
        _conn.execute("PRAGMA synchronous=NORMAL")
    return _conn


def _ensure_schema():
    conn = _db()
    with conn:
        for table, spec in TABLES.items():
            nums = spec["_numset"]
            cols_sql = ", ".join(f"{_q(c)} {'NUMERIC' if c in nums else 'TEXT'}" for c in spec["cols"])
            pk = f", PRIMARY KEY ({', '.join(_q(k) for k in spec['key'])})" if spec["key"] else ""
            conn.execute(f"CREATE TABLE IF NOT EXISTS {table} ({cols_sql}{pk})")
            existing = {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}
            for c in spec["cols"]:
                if c not in existing:
                    conn.execute(f"ALTER TABLE {table} ADD COLUMN {_q(c)} {'NUMERIC' if c in nums else 'TEXT'}")
            for i, cols in enumerate(spec.get("index", [])):
                conn.execute(f"CREATE INDEX IF NOT EXISTS idx_{table}_{i} ON {table} ({', '.join(_q(c) for c in cols)})")
        conn.execute('CREATE INDEX IF NOT EXISTS idx_assets_site ON assets ("현장", "상태")')
        conn.execute("CREATE TABLE IF NOT EXISTS _meta (k TEXT PRIMARY KEY, v TEXT)")


def _sql_value(value, is_num: bool):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return int(value)
    if is_num:
        return to_num(value)
    return value if isinstance(value, str) else clean_str(value)


def _from_row(spec: dict, row: tuple) -> dict:
    rec = dict(zip(spec["cols"], row))
    for c in spec["dates"]:
        v = rec[c]
        if v:
            try:
                rec[c] = date.fromisoformat(v[:10])
            except ValueError:
                rec[c] = None
        else:
            rec[c] = None
    for c in spec["_text"]:
        v = rec[c]
        if v is None:
            rec[c] = ""
        elif not isinstance(v, str):
            rec[c] = clean_str(v)
    return rec


def _normalize(table: str, row: dict) -> dict:
    spec = TABLES[table]
    out = {}
    for col in spec["cols"]:
        value = row.get(col)
        if col in spec["dates"]:
            out[col] = to_date(value)
        elif col in spec["_numset"]:
            out[col] = to_num(value)
        else:
            out[col] = clean_str(value)
    return out


def _key_of(spec: dict, rec: dict):
    return tuple(rec[k] for k in spec["key"]) if len(spec["key"]) > 1 else rec[spec["key"][0]]


def load(table: str) -> list[dict]:
    """캐시된 테이블 전체. 읽기 전용으로 쓰세요(고칠 때는 dict(rec)로 복사한 뒤 upsert)."""
    spec = TABLES[table]
    if not spec["cache"]:
        return query(table)
    with _lock:
        c = _cache.get(table)
        if c is None:
            order = ", ".join(_q(k) for k in spec["key"])
            cur = _db().execute(f"SELECT {', '.join(_q(x) for x in spec['cols'])} FROM {table} ORDER BY {order}")
            rows = [_from_row(spec, r) for r in cur]
            c = {"rows": rows, "index": {_key_of(spec, r): i for i, r in enumerate(rows)}, "ver": 0, "sorted": True}
            _cache[table] = c
        return c["rows"]


def is_loaded(table: str = "assets") -> bool:
    return table in _cache


def data_version(table: str) -> int:
    """표가 바뀔 때마다 커지는 번호(파생 계산 캐시의 키로 사용)."""
    load(table)
    return _ver_bump.get(table, 0) * 1_000_000 + _cache[table]["ver"]


def lookup(table: str, key) -> dict | None:
    load(table)
    c = _cache[table]
    pos = c["index"].get(key)
    return c["rows"][pos] if pos is not None else None


def query(table: str, where: str = "", params: tuple = (), order: str = "", limit: int | None = None) -> list[dict]:
    spec = TABLES[table]
    sql = f"SELECT {', '.join(_q(x) for x in spec['cols'])} FROM {table}"
    if where:
        sql += f" WHERE {where}"
    if order:
        sql += f" ORDER BY {order}"
    if limit:
        sql += f" LIMIT {int(limit)}"
    with _lock:
        return [_from_row(spec, r) for r in _db().execute(sql, params)]


def scalar(sql: str, params: tuple = ()):
    with _lock:
        row = _db().execute(sql, params).fetchone()
    return row[0] if row else None


def _insert_sql(table: str, verb: str) -> str:
    cols = TABLES[table]["cols"]
    return f"{verb} INTO {table} ({', '.join(_q(c) for c in cols)}) VALUES ({', '.join('?' for _ in cols)})"


def upsert(table: str, records: list[dict]):
    """바뀐 행만 저장(INSERT OR REPLACE)하고 캐시에도 바로 반영합니다."""
    if not records:
        return
    spec = TABLES[table]
    recs = [_normalize(table, r) for r in records]
    nums = spec["_numset"]
    with _lock:
        conn = _db()
        try:
            with conn:
                conn.executemany(_insert_sql(table, "INSERT OR REPLACE"),
                                 [tuple(_sql_value(r[c], c in nums) for c in spec["cols"]) for r in recs])
        except sqlite3.OperationalError as exc:
            raise DataError(f"데이터베이스에 저장하지 못했습니다: {exc}. 다른 프로그램이 자산관리.db를 쓰고 있는지 확인하세요.") from exc
        c = _cache.get(table)
        if c is not None and spec["key"]:
            for r in recs:
                k = _key_of(spec, r)
                pos = c["index"].get(k)
                if pos is None:
                    c["index"][k] = len(c["rows"])
                    c["rows"].append(r)
                    c["sorted"] = False
                else:
                    c["rows"][pos] = r
            c["ver"] += 1


def insert_bulk(table: str, records: list[dict], normalized: bool = False):
    """대량 추가(예시 데이터·이력·조사 결과). 캐시는 비우고 다음 조회 때 다시 읽습니다."""
    if not records:
        return
    spec = TABLES[table]
    nums = spec["_numset"]
    recs = records if normalized else [_normalize(table, r) for r in records]
    with _lock:
        conn = _db()
        with conn:
            conn.executemany(_insert_sql(table, "INSERT OR REPLACE" if spec["key"] else "INSERT"),
                             (tuple(_sql_value(r.get(c), c in nums) for c in spec["cols"]) for r in recs))
        _invalidate(table)


def _invalidate(table: str):
    old = _cache.pop(table, None)
    _ver_bump[table] = _ver_bump.get(table, 0) + 1
    if table == "assets":
        _derived.clear()
    del old


def init_storage():
    """처음 실행할 때 DB와 기본 현장·설정을 만들고, 예전 엑셀 데이터가 있으면 옮겨 옵니다."""
    with _lock:
        _ensure_schema()
        if not scalar("SELECT v FROM _meta WHERE k='initialized'"):
            _migrate_from_excel()
            if not load("sites"):
                upsert("sites", [dict(s) for s in DEFAULT_SITES])
            with _db() as conn:
                conn.execute("INSERT OR REPLACE INTO _meta VALUES ('initialized', ?)", (now_text(),))
        current = {r["항목"] for r in load("settings")}
        missing = [k for k in DEFAULT_SETTINGS if k not in current]
        if missing:
            upsert("settings", [{"항목": k, "값": DEFAULT_SETTINGS[k][0], "설명": DEFAULT_SETTINGS[k][1]} for k in missing])


def _migrate_from_excel():
    for table, spec in TABLES.items():
        if not spec.get("legacy"):
            continue
        path = DATA_DIR / spec["legacy"][0]
        if not path.exists():
            continue
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception:
            continue
        try:
            if spec["legacy"][1] not in wb.sheetnames:
                continue
            rows = wb[spec["legacy"][1]].iter_rows(values_only=True)
            header = [clean_str(h) for h in (next(rows, None) or [])]
            recs = [{header[i]: v[i] for i in range(min(len(header), len(v))) if header[i]}
                    for v in rows if v and any(x is not None and clean_str(x) for x in v)]
        finally:
            wb.close()
        if recs:
            insert_bulk(table, recs)


# =========================================================
# 내보내기 / 백업
# =========================================================
_HEADER_FILL = PatternFill("solid", fgColor="5B9BD5")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)


def rows_to_xlsx(sheets: dict[str, tuple[list[str], list[dict]]]) -> bytes:
    """{시트명: (열 목록, 행 목록)} → xlsx. write_only 모드라 수십만 행도 처리합니다."""
    wb = Workbook(write_only=True)
    for name, (cols, rows) in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        for i, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(str(c)) * 2 + 4))
        ws.freeze_panes = "B2"
        header = []
        for c in cols:
            cell = WriteOnlyCell(ws, value=c)
            cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
            cell.alignment = Alignment(horizontal="center")
            header.append(cell)
        ws.append(header)
        for r in rows:
            ws.append([None if r.get(c) == "" else r.get(c) for c in cols])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rows_to_csv(cols: list[str], rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(cols)
    for r in rows:
        w.writerow(["" if r.get(c) is None else r.get(c) for c in cols])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # 엑셀에서 한글이 깨지지 않도록 BOM


def table_rows(table: str) -> list[dict]:
    return load(table) if TABLES[table]["cache"] else query(table)


def table_count(table: str) -> int:
    return scalar(f"SELECT COUNT(*) FROM {table}") or 0


def backup_db_bytes() -> bytes:
    """사용 중인 DB를 안전하게 복사한 파일 내용(그대로 data 폴더에 넣으면 복원됩니다)."""
    with _lock:
        tmp = DATA_DIR / "_backup.db"
        if tmp.exists():
            tmp.unlink()
        dst = sqlite3.connect(tmp)
        try:
            _db().backup(dst)
        finally:
            dst.close()
        data = tmp.read_bytes()
        tmp.unlink()
        return data


# =========================================================
# 설정 / 현장
# =========================================================
def get_settings() -> dict:
    values = {k: v for k, (v, _) in DEFAULT_SETTINGS.items()}
    for row in load("settings"):
        values[row["항목"]] = row["값"]
    return values


def capitalization_threshold() -> int:
    return to_num(get_settings().get("자본화 기준금액"), 1000000) or 0


def save_settings(values: dict):
    rows = []
    for k, v in values.items():
        old = lookup("settings", k) or {}
        rows.append({"항목": k, "값": clean_str(v), "설명": old.get("설명") or DEFAULT_SETTINGS.get(k, ("", ""))[1]})
    upsert("settings", rows)


def get_sites() -> list[dict]:
    return sorted(load("sites"), key=lambda s: (s.get("표시순서") or 999, s["현장명"]))


def site_names() -> list[str]:
    return [s["현장명"] for s in get_sites() if s["현장명"]]


def save_sites(rows: list[dict]):
    """현장 목록 전체를 저장합니다. 보유 자산이 있는 현장은 삭제할 수 없습니다."""
    names = [clean_str(r.get("현장명")) for r in rows]
    if any(not n for n in names):
        raise DataError("현장명이 비어 있는 행이 있습니다.")
    if len(set(names)) != len(names):
        raise DataError("같은 현장명이 두 번 이상 있습니다.")
    removed = set(site_names()) - set(names)
    in_use = {a["현장"] for a in active_assets()}
    blocked = sorted(removed & in_use)
    if blocked:
        raise DataError(f"보유 자산이 있는 현장은 삭제할 수 없습니다: {', '.join(blocked)}. 자산을 먼저 다른 현장으로 이동하세요.")
    with _lock:
        if removed:
            with _db() as conn:
                conn.executemany('DELETE FROM sites WHERE "현장명"=?', [(n,) for n in removed])
            _invalidate("sites")
        upsert("sites", rows)


# =========================================================
# 감가상각
# =========================================================
def account_info(asset: dict) -> dict:
    return ACCOUNTS.get(asset.get("계정과목"), ACCOUNTS["비품"])


def months_used(acquired, as_of: date) -> int:
    acquired = acquired if isinstance(acquired, date) else to_date(acquired)
    if not acquired:
        return 0
    return max(0, (as_of.year - acquired.year) * 12 + (as_of.month - acquired.month) + 1)


def asset_life(asset: dict) -> int:
    life = asset.get("내용연수")
    if isinstance(life, (int, float)) and life:
        return int(life)
    return int(to_num(life, None) or account_info(asset)["life"] or 0)


def asset_method(asset: dict) -> str:
    return clean_str(asset.get("상각방법")) or account_info(asset)["method"]


def accumulated_at(asset: dict, months: int) -> float:
    """취득 후 months개월 동안의 상각누계액.

    - 소모품: 취득 시 전액 비용 처리(누계액 = 취득가액)
    - 정액법: (취득가액 - 잔존가액) × 경과월 / 내용연수월
    - 정률법: 세법 방식. 잔존가치 5%로 상각률 r = 1 - 0.05^(1/n), 연 단위 체감 + 연중 월할
    - 내용연수 종료 시 유형자산은 비망가액 1,000원(또는 잔존가액)만 남깁니다.
    """
    info = account_info(asset)
    cost = asset.get("취득가액")
    cost = float(cost if isinstance(cost, (int, float)) else (to_num(cost, 0) or 0))
    if cost <= 0:
        return 0.0
    life = asset_life(asset)
    if info["cls"] == "소모품" or life <= 0:
        return cost
    if months <= 0:
        return 0.0
    residual = asset.get("잔존가액") or 0
    residual = float(residual if isinstance(residual, (int, float)) else (to_num(residual, 0) or 0))
    floor = max(residual, min(0.0 if info["cls"] == "무형자산" else 1000.0, cost))
    total_months = life * 12
    if months >= total_months:
        return cost - floor
    if asset_method(asset) == "정률":
        rate = 1 - (max(residual, cost * 0.05) / cost) ** (1 / life)
        years, rest = divmod(months, 12)
        return min(cost - floor, cost - cost * (1 - rate) ** years * (1 - rate * rest / 12))
    return min(cost - floor, (cost - residual) * months / total_months)


def depreciation(asset: dict, as_of: date | None = None) -> dict:
    """기준일(as_of) 시점의 경과월, 상각누계액, 장부가액, 해당 월 상각비."""
    as_of = as_of or today()
    disposed = asset.get("처분일자")
    if disposed and disposed < as_of:
        as_of = disposed
    m = months_used(asset.get("취득일자"), as_of)
    accum = accumulated_at(asset, m)
    cost = float(to_num(asset.get("취득가액"), 0) or 0)
    return {"months": m, "accum": accum, "nbv": cost - accum, "monthly": accum - accumulated_at(asset, m - 1)}


def _derived_get(name: str, builder, *extra):
    key = (name, data_version("assets"), *extra)
    value = _derived.get(key)
    if value is None:
        for k in [k for k in _derived if k[0] == name and k[1] != key[1]]:
            _derived.pop(k, None)
        value = builder()
        _derived[key] = value
    return value


def _nbv_map() -> dict:
    return _derived_get("nbv", lambda: {a["태그번호"]: depreciation(a)["nbv"] for a in active_assets()}, today())


def book_value(asset: dict, as_of: date | None = None) -> float:
    if asset.get("상태") in CLOSED_STATUSES:
        return 0.0
    if as_of is None and lookup("assets", asset["태그번호"]) is asset:
        v = _nbv_map().get(asset["태그번호"])
        if v is not None:
            return v
    return depreciation(asset, as_of)["nbv"]


def life_remaining_months(asset: dict) -> int | None:
    life = asset_life(asset)
    if account_info(asset)["cls"] == "소모품" or life <= 0:
        return None
    return life * 12 - months_used(asset.get("취득일자"), today())


def depreciation_schedule(year: int, month: int, site: str | None = None) -> list[dict]:
    """결산월 기준 고정자산 명세(소모품 제외)."""
    return _derived_get("sched", lambda: _build_schedule(year, month, site), year, month, site or "")


def _build_schedule(year: int, month: int, site: str | None) -> list[dict]:
    period_end = date(year, month, 28) + timedelta(days=4)
    period_end = period_end - timedelta(days=period_end.day)
    year_start, prev_end = date(year, 1, 1), date(year - 1, 12, 31)
    rows = []
    for a in load("assets"):
        if site and a["현장"] != site:
            continue
        disposed = a["처분일자"]
        if (disposed and disposed < year_start) or a["계정과목"] == "소모품":
            continue
        acquired = a["취득일자"]
        if not acquired or acquired > period_end:
            continue
        end = depreciation(a, period_end)
        prev = depreciation(a, prev_end) if acquired <= prev_end else {"accum": 0.0, "nbv": 0.0}
        gone = bool(disposed and disposed <= period_end)
        cost = float(a["취득가액"] or 0)
        life = asset_life(a)
        rows.append({
            "태그번호": a["태그번호"], "품명": a["품명"], "계정과목": a["계정과목"], "현장": a["현장"],
            "취득일자": acquired, "취득가액": cost, "내용연수": life, "상각방법": asset_method(a),
            "기초 장부가": prev["nbv"] if acquired <= prev_end else 0.0,
            "당기 취득": cost if acquired >= year_start else 0.0,
            "당기 상각비": end["accum"] - prev["accum"],
            "월 상각비": 0.0 if gone else end["monthly"],
            "상각누계액": 0.0 if gone else end["accum"],
            "기말 장부가": 0.0 if gone else end["nbv"],
            "경과월": min(end["months"], life * 12),
            "비고": f"{a['처분방법'] or '처분'} {disposed}" if gone else "",
        })
    rows.sort(key=lambda r: (ACCOUNT_NAMES.index(r["계정과목"]) if r["계정과목"] in ACCOUNT_NAMES else 99, r["태그번호"]))
    return rows


# =========================================================
# 자산
# =========================================================
def _assets_sorted() -> list[dict]:
    with _lock:
        rows = load("assets")
        c = _cache["assets"]
        if not c["sorted"]:
            rows.sort(key=lambda a: a["태그번호"])
            c["index"] = {a["태그번호"]: i for i, a in enumerate(rows)}
            c["sorted"] = True
        return rows


def active_assets() -> list[dict]:
    """보유 자산(망실·처분완료 제외), 태그번호 순."""
    return _derived_get("active", lambda: [a for a in _assets_sorted() if a["상태"] not in CLOSED_STATUSES])


def closed_assets() -> list[dict]:
    """처분·망실 자산, 처분일자 최신 순."""
    def build():
        rows = [a for a in _assets_sorted() if a["상태"] in CLOSED_STATUSES]
        rows.sort(key=lambda a: (a["처분일자"] or date.min, a["태그번호"]), reverse=True)
        return rows
    return _derived_get("closed", build)


def get_assets(site: str | None = None, include_closed: bool = True) -> list[dict]:
    rows = _assets_sorted() if include_closed else active_assets()
    if site:
        return [a for a in rows if a["현장"] == site]
    return list(rows)


def get_asset(tag: str) -> dict | None:
    return lookup("assets", clean_str(tag).upper())


def is_active(asset: dict) -> bool:
    return asset.get("상태") not in CLOSED_STATUSES


def add_history(tag: str, kind: str, text: str, actor: str):
    add_history_many([(tag, kind, text)], actor)


def add_history_many(items: list[tuple[str, str, str]], actor: str):
    if not items:
        return
    stamp = now_text()
    insert_bulk("history", [{"일시": stamp, "태그번호": t, "구분": k, "내용": x, "처리자": actor} for t, k, x in items], normalized=True)


def get_history(tag: str | None = None, limit: int | None = None) -> list[dict]:
    if tag:
        return query("history", '"태그번호"=?', (tag,), order='"일시" DESC, rowid DESC', limit=limit)
    return query("history", order='"일시" DESC, rowid DESC', limit=limit or 50)


def _tag_counters() -> dict[str, int]:
    """태그 접두어(예: FUR-2026-)별 현재 최대 일련번호."""
    out: dict[str, int] = {}
    for a in load("assets"):
        tag = a["태그번호"]
        p = tag.rfind("-")
        if p > 0 and tag[p + 1:].isdigit():
            prefix, n = tag[:p + 1], int(tag[p + 1:])
            if n > out.get(prefix, 0):
                out[prefix] = n
    return out


def _next_tags(account: str, count: int, year: int | None = None, counters: dict | None = None) -> list[str]:
    prefix = f"{ACCOUNTS.get(account, ACCOUNTS['비품'])['code']}-{year or today().year}-"
    counters = counters if counters is not None else _tag_counters()
    start = counters.get(prefix, 0) + 1
    counters[prefix] = start + count - 1
    return [f"{prefix}{n:04d}" for n in range(start, start + count)]


def _validate_asset_fields(fields: dict):
    if not clean_str(fields.get("품명")):
        raise DataError("품명을 입력하세요.")
    if fields.get("계정과목") not in ACCOUNTS:
        raise DataError("계정과목을 선택하세요.")
    if clean_str(fields.get("현장")) not in site_names():
        raise DataError("현장을 선택하세요. 현장은 설정에서 추가할 수 있습니다.")
    if not to_date(fields.get("취득일자")):
        raise DataError("취득일자를 입력하세요.")
    cost = to_num(fields.get("취득가액"))
    if cost is None or cost < 0:
        raise DataError("취득가액을 0 이상으로 입력하세요.")


def _new_asset(fields: dict, tag: str) -> dict:
    info = ACCOUNTS[fields["계정과목"]]
    rec = {c: "" for c in TABLES["assets"]["cols"]}
    rec.update({k: v for k, v in fields.items() if k in rec})
    rec.update({"태그번호": tag, "자산구분": info["cls"], "내용연수": to_num(fields.get("내용연수")) or info["life"],
                "상각방법": clean_str(fields.get("상각방법")) or info["method"], "잔존가액": to_num(fields.get("잔존가액"), 0) or 0,
                "상태": clean_str(fields.get("상태")) or "사용중", "등록일시": now_text()})
    return rec


def register_assets(fields: dict, qty: int, actor: str, serials: list[str] | None = None,
                    pr_no: str = "", history_text: str = "") -> list[str]:
    _validate_asset_fields(fields)
    qty = max(1, min(int(qty or 1), 500))
    with _lock:
        tags = _next_tags(fields["계정과목"], qty)
        recs = []
        for i, tag in enumerate(tags):
            rec = _new_asset(fields, tag)
            rec["구매요청번호"] = pr_no
            if serials and i < len(serials):
                rec["시리얼번호"] = clean_str(serials[i])
            recs.append(rec)
        upsert("assets", recs)
    add_history_many([(t, "등록", history_text or f"자산대장 등록 · {fields['현장']}") for t in tags], actor)
    return tags


def update_asset(tag: str, changes: dict, actor: str, kind: str = "정보수정", text: str = ""):
    with _lock:
        cur = get_asset(tag)
        if cur is None:
            raise DataError(f"{tag} 자산을 찾을 수 없습니다.")
        rec = dict(cur)
        if "계정과목" in changes and changes["계정과목"] in ACCOUNTS:
            changes = dict(changes, 자산구분=ACCOUNTS[changes["계정과목"]]["cls"])
        diffs = []
        for k, v in changes.items():
            if k not in rec or k == "태그번호":
                continue
            if clean_str(rec[k]) != clean_str(v):
                if k in ("현장", "사용자", "관리부서", "설치위치", "상태", "취득가액", "계정과목"):
                    diffs.append(f"{k} {clean_str(rec[k]) or '-'} → {clean_str(v) or '-'}")
                rec[k] = v
        upsert("assets", [rec])
    if diffs or text:
        add_history(tag, kind, text or ", ".join(diffs), actor)


def update_assets_bulk(changes_by_tag: dict[str, dict], actor: str, kind: str, text_by_tag: dict[str, str]):
    with _lock:
        recs = []
        for tag, changes in changes_by_tag.items():
            cur = get_asset(tag)
            if cur:
                rec = dict(cur)
                rec.update({k: v for k, v in changes.items() if k in rec})
                recs.append(rec)
        upsert("assets", recs)
    add_history_many([(t, kind, text_by_tag.get(t, "")) for t in changes_by_tag], actor)


def handover(tag: str, dept: str, holder: str, location: str, reason: str, actor: str):
    a = get_asset(tag)
    if not a:
        raise DataError("자산을 찾을 수 없습니다.")
    text = f"{reason}: {a['사용자'] or '-'}({a['관리부서'] or '-'}) → {holder or '-'}({dept or '-'})"
    if location and location != a["설치위치"]:
        text += f" · 위치 {a['설치위치'] or '-'} → {location}"
    update_asset(tag, {"관리부서": dept, "사용자": holder, "설치위치": location}, actor, "인수인계", text)


IMPORT_TEMPLATE_COLS = ["품명", "규격", "시리얼번호", "계정과목", "현장", "관리부서", "사용자", "설치위치",
                        "취득일자", "취득가액", "내용연수", "상각방법", "잔존가액", "공급업체", "상태", "비고"]


def import_assets_from_rows(rows: list[dict], actor: str) -> tuple[list[str], list[str]]:
    """엑셀 일괄 등록. 오류가 하나라도 있으면 아무것도 등록하지 않고 오류 목록을 돌려줍니다."""
    errors, prepared = [], []
    sites = set(site_names())
    for i, r in enumerate(rows, start=2):
        name = clean_str(r.get("품명"))
        if not name:
            continue
        acc, site = clean_str(r.get("계정과목")), clean_str(r.get("현장"))
        acquired, cost = to_date(r.get("취득일자")), to_num(r.get("취득가액"))
        problems = [label for ok, label in ((acc in ACCOUNTS, f"계정과목 '{acc}'"), (site in sites, f"현장 '{site}'"),
                                            (bool(acquired), "취득일자"), (cost is not None, "취득가액")) if not ok]
        if problems:
            errors.append(f"{i}행 {name}: {', '.join(problems)} 확인 필요")
            continue
        status = clean_str(r.get("상태"))
        prepared.append({"품명": name, "규격": clean_str(r.get("규격")), "시리얼번호": clean_str(r.get("시리얼번호")),
                         "계정과목": acc, "현장": site, "관리부서": clean_str(r.get("관리부서")), "사용자": clean_str(r.get("사용자")),
                         "설치위치": clean_str(r.get("설치위치")), "취득일자": acquired, "취득가액": cost,
                         "내용연수": to_num(r.get("내용연수")), "상각방법": clean_str(r.get("상각방법")),
                         "잔존가액": to_num(r.get("잔존가액"), 0), "공급업체": clean_str(r.get("공급업체")),
                         "상태": status if status in ASSET_STATUSES and status not in CLOSED_STATUSES else "사용중",
                         "비고": clean_str(r.get("비고"))})
    if errors:
        return [], errors
    with _lock:
        counters = _tag_counters()
        recs = [_new_asset(f, _next_tags(f["계정과목"], 1, f["취득일자"].year, counters)[0]) for f in prepared]
        upsert("assets", recs)
    tags = [r["태그번호"] for r in recs]
    add_history_many([(t, "등록", "엑셀 일괄 등록") for t in tags], actor)
    return tags, []


def read_upload_rows(file_obj) -> list[dict]:
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        header = [clean_str(h) for h in (next(rows, None) or [])]
        return [{header[i]: v[i] for i in range(min(len(header), len(v))) if header[i]}
                for v in rows if v and any(x is not None for x in v)]
    finally:
        wb.close()


# =========================================================
# 요청 공통
# =========================================================
def get_requests(kind: str) -> list[dict]:
    return sorted(load(kind), key=lambda r: r["요청일시"], reverse=True)


def get_request(kind: str, no: str) -> dict | None:
    return lookup(kind, no)


def is_open(kind: str, rec: dict) -> bool:
    return rec["상태"] not in (FLOWS[kind][-1], REJECTED)


def _next_no(kind: str) -> str:
    prefix = f"{KIND_PREFIX[kind]}-{today():%y%m}-"
    used = [int(r["요청번호"][len(prefix):]) for r in load(kind)
            if r["요청번호"].startswith(prefix) and r["요청번호"][len(prefix):].isdigit()]
    return f"{prefix}{max(used, default=0) + 1:03d}"


def _log_line(action: str, actor: str, note: str = "") -> str:
    return f"[{now_text()}] {action} · {actor}" + (f" · {note}" if note else "")


def _append_log(rec: dict, action: str, actor: str, note: str = ""):
    rec["처리이력"] = (rec.get("처리이력") + "\n" if rec.get("처리이력") else "") + _log_line(action, actor, note)


def _update_request(kind: str, no: str, expect: list[str] | None, mutate):
    """요청 1건을 최신 상태로 확인한 뒤 고쳐 저장합니다."""
    with _lock:
        cur = lookup(kind, no)
        if cur is None:
            raise DataError(f"{no} 요청을 찾을 수 없습니다.")
        if expect and cur["상태"] not in expect:
            raise DataError(f"{no}은(는) 이미 '{cur['상태']}' 상태입니다. 화면을 새로고침하세요.")
        rec = dict(cur)
        result = mutate(rec)
        upsert(kind, [rec])
        return result if result is not None else rec


def _tags_of(rec: dict) -> list[str]:
    return [t.strip() for t in clean_str(rec.get("태그번호")).split(",") if t.strip()]


def request_assets(rec: dict) -> list[dict]:
    return [a for a in (get_asset(t) for t in _tags_of(rec)) if a]


def _new_request(kind: str, fields: dict) -> dict:
    rec = {c: "" for c in TABLES[kind]["cols"]}
    rec.update({k: v for k, v in fields.items() if k in rec})
    return rec


# =========================================================
# 1. 구매요청 → 견적 → 발주 → 검수 → 등록
# =========================================================
def create_purchase_request(fields: dict, actor: str) -> str:
    for key, label in (("품명", "품명"), ("구매사유", "구매 사유")):
        if not clean_str(fields.get(key)):
            raise DataError(f"{label}을(를) 입력하세요.")
    if (to_num(fields.get("수량")) or 0) < 1:
        raise DataError("수량은 1 이상이어야 합니다.")
    if fields.get("현장") not in site_names():
        raise DataError("요청 현장을 선택하세요.")
    with _lock:
        no = _next_no("pr")
        rec = _new_request("pr", fields)
        rec.update({"요청번호": no, "상태": "요청", "요청일시": now_text(), "요청자": actor})
        _append_log(rec, "구매요청", actor, clean_str(fields.get("구매사유"))[:60])
        upsert("pr", [rec])
    return no


def pr_decide(no: str, approve: bool, note: str, actor: str):
    if not approve and not clean_str(note):
        raise DataError("반려 사유를 입력하세요.")

    def mutate(rec):
        rec["상태"] = "승인" if approve else REJECTED
        rec["검토의견"] = note
        _append_log(rec, "타당성 승인" if approve else "반려", actor, note)
    _update_request("pr", no, ["요청"], mutate)


def pr_quotes(rec: dict) -> list[dict]:
    out = []
    for i in (1, 2, 3):
        vendor = clean_str(rec.get(f"견적{i} 업체"))
        if vendor:
            out.append({"slot": i, "업체": vendor, "단가": to_num(rec.get(f"견적{i} 단가"), 0) or 0,
                        "조건": clean_str(rec.get(f"견적{i} 조건"))})
    return out


def pr_save_quotes(no: str, quotes: list[dict], actor: str):
    valid = [q for q in quotes if clean_str(q.get("업체"))]
    if any((to_num(q.get("단가")) or 0) <= 0 for q in valid):
        raise DataError("견적 단가를 입력하세요.")

    def mutate(rec):
        for i in (1, 2, 3):
            q = valid[i - 1] if i <= len(valid) else {}
            rec[f"견적{i} 업체"] = clean_str(q.get("업체"))
            rec[f"견적{i} 단가"] = to_num(q.get("단가"))
            rec[f"견적{i} 조건"] = clean_str(q.get("조건"))
        _append_log(rec, "견적 등록", actor, ", ".join(f"{q['업체']} {int(to_num(q['단가'])):,}원" for q in valid))
    _update_request("pr", no, ["승인"], mutate)


def pr_choose_vendor(no: str, slot: int, reason: str, actor: str):
    def mutate(rec):
        quotes = pr_quotes(rec)
        if len(quotes) < 2:
            raise DataError("비교견적이 2건 이상 있어야 업체를 선정할 수 있습니다.")
        chosen = next((q for q in quotes if q["slot"] == slot), None)
        if not chosen:
            raise DataError("선정할 견적을 고르세요.")
        if not clean_str(reason):
            raise DataError("선정 사유를 입력하세요.")
        rec.update({"선정업체": chosen["업체"], "선정단가": chosen["단가"], "선정사유": reason, "상태": "발주"})
        _append_log(rec, "업체 선정·발주", actor, f"{chosen['업체']} {int(chosen['단가']):,}원 · {reason}")
    _update_request("pr", no, ["승인"], mutate)


def pr_inspect(no: str, spec_ok: bool, received_qty: int, no_damage: bool, serials: str, note: str, actor: str) -> bool:
    def mutate(rec):
        qty_ok = int(received_qty or 0) == int(to_num(rec["수량"], 0) or 0)
        passed = spec_ok and no_damage and qty_ok
        rec.update({"검수 규격": "적합" if spec_ok else "부적합", "검수 입고수량": int(received_qty or 0),
                    "검수 파손": "이상 없음" if no_damage else "파손 있음", "검수 시리얼": clean_str(serials),
                    "검수 메모": note, "검수자": actor, "검수일시": now_text(), "상태": "검수완료" if passed else "발주"})
        fails = [x for x, ok in (("규격", spec_ok), ("수량", qty_ok), ("파손", no_damage)) if not ok]
        _append_log(rec, "검수 합격" if passed else "검수 부적합(재입고 후 재검수)", actor, note or ", ".join(fails))
        return passed
    return _update_request("pr", no, ["발주"], mutate)


def pr_register(no: str, asset_fields: dict, actor: str) -> list[str]:
    rec = get_request("pr", no)
    if not rec or rec["상태"] != "검수완료":
        raise DataError("검수가 끝난 요청만 자산대장에 등록할 수 있습니다.")
    serials = [s.strip() for s in clean_str(rec["검수 시리얼"]).split(",") if s.strip()]
    tags = register_assets(asset_fields, int(to_num(rec["수량"], 1) or 1), actor, serials=serials, pr_no=no,
                           history_text=f"구매요청 {no} · {rec['선정업체']} 입고 검수 후 등록")

    def mutate(r):
        r["상태"] = "등록완료"
        r["등록 태그번호"] = ", ".join(tags)
        _append_log(r, "자산 등록", actor, ", ".join(tags))
    _update_request("pr", no, ["검수완료"], mutate)
    return tags


# =========================================================
# 2. 이동신청 (보내는 현장 승인 → 받는 현장 인수확인)
# =========================================================
def create_transfer(tags: list[str], to_site: str, to_dept: str, to_holder: str, reason: str, actor: str) -> list[str]:
    if not tags:
        raise DataError("이동할 자산을 고르세요.")
    if to_site not in site_names():
        raise DataError("받는 현장을 선택하세요.")
    if not clean_str(reason):
        raise DataError("신청 사유를 입력하세요.")
    open_tags = {t for r in load("move") if is_open("move", r) for t in _tags_of(r)}
    groups: dict[str, list[dict]] = {}
    for t in tags:
        a = get_asset(t)
        if not a or not is_active(a) or a["상태"] == "처분대기":
            raise DataError(f"{t}은(는) 이동할 수 없는 상태입니다.")
        if t in open_tags:
            raise DataError(f"{t}은(는) 이미 진행 중인 이동신청이 있습니다.")
        if a["현장"] == to_site:
            raise DataError(f"{t}은(는) 이미 {to_site} 자산입니다. 같은 현장 안의 변경은 '인수인계'를 사용하세요.")
        groups.setdefault(a["현장"], []).append(a)
    created = []
    with _lock:
        for from_site, items in groups.items():
            no = _next_no("move")
            rec = _new_request("move", {"요청번호": no, "상태": "요청", "요청일시": now_text(), "요청자": actor,
                                        "태그번호": ", ".join(a["태그번호"] for a in items),
                                        "품명": ", ".join(sorted({a["품명"] for a in items})),
                                        "보내는 현장": from_site, "받는 현장": to_site, "받는 부서": to_dept,
                                        "새 사용자": to_holder, "사유": reason})
            _append_log(rec, "이동신청", actor, reason)
            upsert("move", [rec])
            created.append(no)
    return created


def transfer_decide(no: str, approve: bool, note: str, actor: str):
    if not approve and not clean_str(note):
        raise DataError("반려 사유를 입력하세요.")

    def mutate(rec):
        rec["상태"] = "승인" if approve else REJECTED
        rec["승인자"], rec["승인일시"] = actor, now_text()
        _append_log(rec, "보내기 승인" if approve else "반려", actor, note)
    _update_request("move", no, ["요청"], mutate)


def transfer_receive(no: str, note: str, actor: str):
    rec = get_request("move", no)
    if not rec or rec["상태"] != "승인":
        raise DataError("승인된 이동신청만 인수확인할 수 있습니다.")
    changes, texts = {}, {}
    for a in request_assets(rec):
        changes[a["태그번호"]] = {"현장": rec["받는 현장"], "관리부서": rec["받는 부서"] or a["관리부서"],
                                 "사용자": rec["새 사용자"], "상태": "사용중" if a["상태"] == "유휴" else a["상태"]}
        texts[a["태그번호"]] = f"{rec['보내는 현장']} → {rec['받는 현장']} ({no})" + (f" · 사용자 {rec['새 사용자']}" if rec["새 사용자"] else "")
    update_assets_bulk(changes, actor, "현장 이동", texts)

    def mutate(r):
        r["상태"] = "완료"
        r["인수자"], r["인수일시"] = actor, now_text()
        _append_log(r, "인수 확인", actor, note)
    _update_request("move", no, ["승인"], mutate)


# =========================================================
# 3. 분실·도난·파손
# =========================================================
def create_incident(tag: str, kind: str, occurred, place: str, person: str, statement: str,
                    police: str, actor: str) -> str:
    a = get_asset(tag)
    if not a or not is_active(a):
        raise DataError("신고할 수 있는 자산이 아닙니다.")
    if not clean_str(statement):
        raise DataError("경위를 입력하세요.")
    with _lock:
        no = _next_no("inc")
        rec = _new_request("inc", {"요청번호": no, "상태": "접수", "요청일시": now_text(), "신고자": actor,
                                   "태그번호": a["태그번호"], "품명": a["품명"], "현장": a["현장"], "구분": kind,
                                   "발생일": to_date(occurred) or today(), "발생장소": place, "관련자": person,
                                   "경위": statement, "경찰신고번호": police})
        _append_log(rec, f"{kind} 경위서 접수", actor)
        upsert("inc", [rec])
    changes = {"상태": "수리중", "이전상태": a["상태"]} if kind == "파손" else {}
    update_asset(a["태그번호"], changes, actor, "사고 접수", f"{kind} · {statement[:60]} ({no})")
    return no


def incident_investigate(no: str, fields: dict, actor: str):
    if not clean_str(fields.get("조사의견")):
        raise DataError("조사 의견을 입력하세요.")

    def mutate(rec):
        for k in ("고의과실", "변상여부", "변상금액", "처리방법", "수리견적", "보험청구", "조사의견"):
            rec[k] = fields.get(k, rec[k])
        rec.update({"상태": "조사완료", "조사자": actor, "조사일시": now_text()})
        _append_log(rec, "조사 완료", actor, f"{rec['고의과실']} · {rec['변상여부']} · {rec['처리방법']}")
        return rec
    rec = _update_request("inc", no, ["접수"], mutate)
    if to_num(fields.get("수리견적")):
        update_asset(rec["태그번호"], {"수리견적": to_num(fields["수리견적"])}, actor, "수리 견적",
                     f"수리 견적 {int(to_num(fields['수리견적'])):,}원")


def incident_close(no: str, actor: str) -> str:
    rec = get_request("inc", no)
    if not rec or rec["상태"] != "조사완료":
        raise DataError("조사가 끝난 건만 종결할 수 있습니다.")
    a = get_asset(rec["태그번호"])
    action = rec["처리방법"]
    message = ""
    if a and "망실" in action:
        nbv = book_value(a)
        comp = float(to_num(rec["변상금액"], 0) or 0)
        update_asset(a["태그번호"], {"상태": "망실", "처분일자": today(), "처분방법": "망실", "처분처": "",
                                   "처분금액": comp, "처분손익": round(comp - nbv), "처분요청번호": no},
                     actor, "망실 제적", f"{rec['구분']} · {action} · 손실 {round(nbv - comp):,}원 ({no})")
        message = f"망실 처리했습니다. 장부가 {round(nbv):,}원 중 {round(nbv - comp):,}원이 손실로 인식됩니다."
    elif a:
        restore = a["이전상태"] if a["이전상태"] and a["이전상태"] not in CLOSED_STATUSES else "사용중"
        update_asset(a["태그번호"], {"상태": restore, "이전상태": ""}, actor, "사고 종결", f"{action} ({no})")
        message = f"{a['태그번호']}을(를) '{restore}' 상태로 되돌렸습니다."

    def mutate(r):
        r["상태"] = "종결"
        r["종결일시"] = now_text()
        _append_log(r, "종결", actor, action)
    _update_request("inc", no, ["조사완료"], mutate)
    return message


# =========================================================
# 4. 처분 (신청 → 위원회 승인 → 집행 → 제적)
# =========================================================
def disposal_candidates(site: str | None = None) -> list[dict]:
    out = []
    for a in active_assets():
        if site and a["현장"] != site:
            continue
        if a["상태"] not in ("사용중", "유휴", "수리중"):
            continue
        reasons = []
        left = life_remaining_months(a)
        if left is not None and left <= 0:
            reasons.append("내용연수 만료")
        repair = a["수리견적"] or 0
        if repair and repair > book_value(a):
            reasons.append(f"수리비 {int(repair):,}원 > 장부가")
        if a["상태"] == "유휴" and left is not None and 0 < left <= 12:
            reasons.append("유휴 · 1년 내 만료")
        if reasons:
            out.append({**a, "장부가액": book_value(a), "검토사유": " / ".join(reasons)})
    return out


def needs_secure_wipe(assets: list[dict]) -> bool:
    return any(STORAGE_DEVICE_PATTERN.search(f"{a['품명']} {a['규격']}") for a in assets)


def create_disposal(tags: list[str], reason: str, method: str, detail: str, actor: str) -> list[str]:
    if not tags:
        raise DataError("처분할 자산을 고르세요.")
    if not clean_str(detail):
        raise DataError("상세 사유를 입력하세요.")
    groups: dict[str, list[dict]] = {}
    for t in tags:
        a = get_asset(t)
        if not a or a["상태"] not in ("사용중", "유휴", "수리중"):
            raise DataError(f"{t}은(는) 처분 신청할 수 없는 상태입니다.")
        groups.setdefault(a["현장"], []).append(a)
    created = []
    with _lock:
        for site, items in groups.items():
            no = _next_no("dsp")
            rec = _new_request("dsp", {"요청번호": no, "상태": "신청", "요청일시": now_text(), "신청자": actor, "현장": site,
                                       "태그번호": ", ".join(a["태그번호"] for a in items),
                                       "품명": ", ".join(sorted({a["품명"] for a in items})),
                                       "불용사유": reason, "희망 처분방법": method, "상세사유": detail,
                                       "장부가 합계": round(sum(book_value(a) for a in items))})
            _append_log(rec, "처분 신청", actor, f"{reason} · {method}")
            upsert("dsp", [rec])
            created.append(no)
            update_assets_bulk({a["태그번호"]: {"상태": "처분대기", "이전상태": a["상태"]} for a in items}, actor,
                               "처분 신청", {a["태그번호"]: f"{reason} · {method} ({no})" for a in items})
    return created


def disposal_decide(no: str, approve: bool, note: str, actor: str):
    if not clean_str(note):
        raise DataError("위원회 의결 내용 또는 반려 사유를 입력하세요.")

    def mutate(rec):
        rec["상태"] = "승인" if approve else REJECTED
        rec["위원회 의견"], rec["승인일시"] = note, now_text()
        _append_log(rec, "불용품 판정 승인" if approve else "반려", actor, note)
        return rec
    rec = _update_request("dsp", no, ["신청"], mutate)
    if not approve:
        items = request_assets(rec)
        update_assets_bulk({a["태그번호"]: {"상태": a["이전상태"] or "사용중", "이전상태": ""} for a in items},
                           actor, "처분 반려", {a["태그번호"]: note for a in items})


def disposal_execute(no: str, method: str, exec_date, vendor: str, proceeds, wiped: bool, cert: str,
                     note: str, actor: str):
    rec = get_request("dsp", no)
    if not rec:
        raise DataError("요청을 찾을 수 없습니다.")
    if not clean_str(vendor):
        raise DataError("폐기업체·매수자·기증처를 입력하세요.")
    if needs_secure_wipe(request_assets(rec)) and (not wiped or not clean_str(cert)):
        raise DataError("PC·저장장치가 포함되어 있습니다. 데이터 완전 삭제(보안 파기) 확인과 파기증명서 번호가 필요합니다.")

    def mutate(r):
        r.update({"처분방법": method, "집행일": to_date(exec_date) or today(), "처분처": vendor,
                  "처분금액": to_num(proceeds, 0) or 0, "보안파기 확인": "확인" if wiped else "",
                  "파기증명번호": cert, "상태": "집행완료"})
        _append_log(r, "처분 집행", actor, f"{method} · {vendor} · {int(to_num(proceeds, 0) or 0):,}원" + (f" · {note}" if note else ""))
    _update_request("dsp", no, ["승인"], mutate)


def disposal_preview(rec: dict) -> dict:
    items = request_assets(rec)
    when = to_date(rec.get("집행일")) or today()
    nbvs = [depreciation(a, when)["nbv"] for a in items]
    total = sum(nbvs)
    proceeds = float(to_num(rec.get("처분금액"), 0) or 0)
    return {"items": items, "nbvs": nbvs, "total_nbv": total, "proceeds": proceeds, "gain_loss": proceeds - total}


def disposal_derecognize(no: str, actor: str) -> float:
    rec = get_request("dsp", no)
    if not rec or rec["상태"] != "집행완료":
        raise DataError("집행이 끝난 처분만 제적할 수 있습니다.")
    p = disposal_preview(rec)
    changes, texts = {}, {}
    for a, nbv in zip(p["items"], p["nbvs"]):
        share = p["proceeds"] * (nbv / p["total_nbv"]) if p["total_nbv"] else p["proceeds"] / max(1, len(p["items"]))
        changes[a["태그번호"]] = {"상태": "처분완료", "처분일자": to_date(rec["집행일"]), "처분방법": rec["처분방법"],
                                 "처분처": rec["처분처"], "처분금액": round(share), "처분손익": round(share - nbv),
                                 "처분요청번호": no, "이전상태": ""}
        texts[a["태그번호"]] = f"{rec['처분방법']} · {rec['처분처']} · 처분손익 {round(share - nbv):,}원 ({no})"
    update_assets_bulk(changes, actor, "처분 제적", texts)

    def mutate(r):
        r.update({"상태": "제적완료", "장부가 합계": round(p["total_nbv"]), "처분손익": round(p["gain_loss"]), "제적일시": now_text()})
        _append_log(r, "장부 제적", actor, f"처분손익 {round(p['gain_loss']):,}원")
    _update_request("dsp", no, ["집행완료"], mutate)
    return p["gain_loss"]


# =========================================================
# 재물조사
# =========================================================
def get_audits() -> list[dict]:
    return sorted(load("audits"), key=lambda r: r["시작일시"], reverse=True)


def open_audit() -> dict | None:
    return next((a for a in get_audits() if a["상태"] == "진행중"), None)


def get_audit_items(audit_no: str, site: str | None = None) -> list[dict]:
    if site:
        return query("audit_items", '"회차번호"=? AND "현장"=?', (audit_no, site), order='"현장", "태그번호"')
    return query("audit_items", '"회차번호"=?', (audit_no,), order='"현장", "태그번호"')


def audit_stats(items: list[dict]) -> dict:
    found = sum(1 for r in items if r["결과"] == "확인")
    missing = sum(1 for r in items if r["결과"] == "미발견")
    return {"total": len(items), "found": found, "missing": missing, "todo": len(items) - found - missing,
            "idle": sum(1 for r in items if r["유휴"] == "Y")}


def _audit_counts(audit_no: str) -> dict:
    with _lock:
        row = _db().execute(
            'SELECT COUNT(*), SUM("결과"=\'확인\'), SUM("결과"=\'미발견\'), SUM("유휴"=\'Y\') FROM audit_items WHERE "회차번호"=?',
            (audit_no,)).fetchone()
    total, found, missing, idle = (row[0] or 0, row[1] or 0, row[2] or 0, row[3] or 0)
    return {"total": total, "found": found, "missing": missing, "todo": total - found - missing, "idle": idle}


def start_audit(name: str, base_date, scope: str, actor: str) -> str:
    if open_audit():
        raise DataError("진행 중인 재물조사가 있습니다. 마감한 뒤 새 조사를 시작하세요.")
    targets = [a for a in active_assets() if scope == "전체" or a["현장"] == scope]
    with _lock:
        no = f"AU-{today():%y%m}-{len(load('audits')) + 1:02d}"
        upsert("audits", [{"회차번호": no, "회차명": name or f"{today().year} 재물조사", "기준일": to_date(base_date) or today(),
                           "대상": scope, "상태": "진행중", "시작자": actor, "시작일시": now_text(), "장부수량": len(targets),
                           "확인": 0, "미발견": 0, "유휴": 0}])
        insert_bulk("audit_items", [{"회차번호": no, "태그번호": a["태그번호"], "품명": a["품명"], "현장": a["현장"],
                                     "설치위치": a["설치위치"] or a["관리부서"], "결과": "미확인", "유휴": "", "메모": "",
                                     "확인자": "", "확인일시": ""} for a in targets], normalized=True)
    return no


def audit_mark(audit_no: str, updates: dict[str, dict], actor: str) -> int:
    """updates: {태그번호: {"결과": 확인|미발견|미확인, "유휴": bool, "메모": str}}"""
    with _lock:
        au = lookup("audits", audit_no)
        if not au or au["상태"] != "진행중":
            raise DataError("진행 중인 재물조사가 아닙니다.")
        stamp = now_text()
        rows = []
        for tag, u in updates.items():
            cur = query("audit_items", '"회차번호"=? AND "태그번호"=?', (audit_no, tag))
            if not cur:
                continue
            r = cur[0]
            new = {"결과": u.get("결과", r["결과"]), "유휴": "Y" if u.get("유휴", r["유휴"] == "Y") else "",
                   "메모": clean_str(u.get("메모", r["메모"]))}
            if any(clean_str(r[k]) != clean_str(v) for k, v in new.items()):
                r.update(new)
                r["확인자"], r["확인일시"] = actor, stamp
                rows.append(r)
        if rows:
            insert_bulk("audit_items", rows, normalized=True)
            s = _audit_counts(audit_no)
            upsert("audits", [dict(au, 확인=s["found"], 미발견=s["missing"], 유휴=s["idle"])])
        return len(rows)


def audit_scan(audit_no: str, code: str, actor: str, allowed_site: str | None = None) -> tuple[bool, str]:
    code = clean_str(code)
    if "tag=" in code:  # QR에 주소가 담긴 경우
        code = code.split("tag=")[-1].split("&")[0]
    code = code.upper()
    found = query("audit_items", '"회차번호"=? AND "태그번호"=?', (audit_no, code))
    if not found:
        asset = get_asset(code)
        if not asset:
            return False, f"{code}: 자산대장에 없는 태그입니다."
        return False, f"{code}: 이번 조사 대상이 아닙니다 ({asset['현장']})."
    item = found[0]
    if allowed_site and item["현장"] != allowed_site:
        return False, f"{code}: {item['현장']} 자산입니다. 담당 현장 자산만 확인할 수 있습니다."
    audit_mark(audit_no, {item["태그번호"]: {"결과": "확인"}}, actor)
    return True, f"{item['태그번호']} {item['품명']} 확인"


def close_audit(audit_no: str, actor: str) -> dict:
    items = get_audit_items(audit_no)
    stats = audit_stats(items)
    name = (lookup("audits", audit_no) or {}).get("회차명", audit_no)
    changes, texts, hist = {}, {}, []
    for r in items:
        a = get_asset(r["태그번호"])
        if not a:
            continue
        if r["유휴"] == "Y" and a["상태"] == "사용중":
            changes[a["태그번호"]] = {"상태": "유휴"}
            texts[a["태그번호"]] = f"{name}: 유휴 판정 · 재배치 대상"
        elif r["결과"] == "미발견":
            hist.append((a["태그번호"], "재물조사", f"{name}: 실물 미발견 · 분실 경위 확인 필요"))
        elif r["결과"] == "확인":
            hist.append((a["태그번호"], "재물조사", f"{name}: 실물 확인"))
    if changes:
        update_assets_bulk(changes, actor, "재물조사", texts)
    add_history_many(hist, actor)
    au = dict(lookup("audits", audit_no))
    au.update({"상태": "완료", "마감일시": now_text(), "확인": stats["found"], "미발견": stats["missing"], "유휴": stats["idle"]})
    upsert("audits", [au])
    return stats


# =========================================================
# 결산 마감
# =========================================================
def get_closings() -> list[dict]:
    return sorted(load("closings"), key=lambda r: r["결산월"], reverse=True)


def closing_items(period: str) -> list[dict]:
    return query("closing_items", '"결산월"=?', (period,), order='"태그번호"')


def close_period(year: int, month: int, actor: str) -> dict:
    key = f"{year}-{month:02d}"
    if lookup("closings", key):
        raise DataError(f"{key}은(는) 이미 마감되었습니다.")
    rows = depreciation_schedule(year, month)
    summary = {"결산월": key, "마감일시": now_text(), "마감자": actor, "자산수": len(rows),
               "취득가액": round(sum(r["취득가액"] for r in rows if r["기말 장부가"] or r["상각누계액"])),
               "당기상각비": round(sum(r["당기 상각비"] for r in rows)),
               "상각누계액": round(sum(r["상각누계액"] for r in rows)),
               "장부가액": round(sum(r["기말 장부가"] for r in rows))}
    with _lock:
        insert_bulk("closing_items", [{**{k: (round(v) if isinstance(v, float) else v) for k, v in r.items()}, "결산월": key} for r in rows])
        upsert("closings", [summary])
    return summary


# =========================================================
# 법인차량 관리 프로그램 연동
# =========================================================
VEHICLE_COLUMNS = ["차량번호", "차종", "구분", "현장", "사용자", "사용자 연락처", "관리자", "관리자 연락처", "상태",
                   "누적주행거리", "차량 구매일", "최근정비일", "정비 시작일", "정비 예정 완료일", "타이어 교체 km수",
                   "엔진오일, 필터 교체 km수", "타이어 교체(4만 km 마다)", "엔진오일, 필터 교환(1만 km 마다)",
                   "차량 교체 4년 마다", "차량 구매 비용"]


def vehicle_template_bytes() -> bytes:
    """법인차량 관리 프로그램과 같은 열 구성의 차량목록.xlsx 샘플 양식."""
    sample = [
        ["12가3456", "그랜저", "구매", "본사", "김대표", "010-1234-5678", "김재열", "010-0000-0000", "운행가능", 48210,
         date(2024, 3, 15), date(2026, 6, 1), None, None, 40000, 50000, "정상", "정상", "정상", "40,500,000원"],
        ["34나5678", "카니발", "구매", "평택", "이전무", "010-2345-6789", "김재열", "010-0000-0000", "정비중", 103540,
         date(2022, 11, 2), date(2026, 5, 20), date(2026, 10, 10), date(2026, 10, 20), 80000, 100000, "정상", "교체 대상", "정상", "43,500,000원"],
        ["56다7890", "쏘렌토", "렌트", "동탄", None, None, "김재열", "010-0000-0000", "운행가능", 35120,
         date(2025, 1, 10), date(2026, 7, 3), None, None, 0, 30000, "정상", "정상", "정상", None],
        ["78라1234", "스타리아", "렌트", "탕정", None, None, "김재열", "010-0000-0000", "운행가능", 61800,
         date(2024, 6, 1), date(2026, 4, 11), None, None, 40000, 60000, "정상", "정상", "정상", None],
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"  # 법인차량 프로그램은 첫 시트를 읽습니다
    ws.append(VEHICLE_COLUMNS)
    for row in sample:
        ws.append(row)
    for i, col in enumerate(VEHICLE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(26, len(col) * 2 + 2))
        fmt = ("yyyy-mm-dd" if col in ("차량 구매일", "최근정비일", "정비 시작일", "정비 예정 완료일") else
               "#,##0" if col in ("누적주행거리", "타이어 교체 km수", "엔진오일, 필터 교체 km수") else None)
        if fmt:
            for r in range(2, len(sample) + 2):
                ws.cell(row=r, column=i).number_format = fmt
    ws.freeze_panes = "B2"
    guide = wb.create_sheet("작성안내")
    guide.append(["열", "작성 방법"])
    for row in [
        ["차량번호", "필수. 예: 12가3456 (자산관리에서는 시리얼번호·연결차량번호로 사용)"],
        ["구분", "구매 = 임원용차량(회사 소유 → 차량운반구 자산으로 등록) / 렌트 = 업무용차량(자산 아님, 연동 제외)"],
        ["현장", "본사, 평택, 동탄, 탕정 등. 자산관리에 없는 현장은 연동할 때 자동으로 추가됩니다"],
        ["상태", "운행가능 또는 정비중 (정비중 → 자산 상태 '수리중')"],
        ["차량 구매일", "날짜. 자산의 취득일자가 됩니다"],
        ["차량 구매 비용", "예: 40,500,000원 또는 40500000. 자산의 취득가액이 됩니다"],
        ["사용자", "임원 차량 사용자. 자산의 사용자로 들어갑니다"],
        ["나머지 열", "법인차량 관리 프로그램에서 쓰는 정비·주행 정보입니다. 비워도 연동에는 문제없습니다"],
    ]:
        guide.append(row)
    guide.column_dimensions["A"].width = 16
    guide.column_dimensions["B"].width = 90
    for c in guide[1]:
        c.fill, c.font = _HEADER_FILL, _HEADER_FONT
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_vehicle_file(source) -> list[dict]:
    if isinstance(source, (str, Path)):
        path = Path(clean_str(source).strip('"'))
        if not path.exists():
            raise DataError(f"차량목록 파일을 찾을 수 없습니다: {path}")
        source = path
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        header = [clean_str(h) for h in (next(rows, None) or [])]
        if "차량번호" not in header:
            raise DataError("차량목록 파일에 '차량번호' 열이 없습니다. 샘플 양식의 머리글을 그대로 사용하세요.")
        i_no = header.index("차량번호")
        return [{header[i]: v[i] for i in range(min(len(header), len(v))) if header[i]}
                for v in rows if v and len(v) > i_no and clean_str(v[i_no])]
    finally:
        wb.close()


def sync_vehicles(vehicles: list[dict], actor: str) -> dict:
    """구매 차량을 '차량운반구' 자산으로 등록·갱신합니다. 렌트 차량은 회사 소유가 아니라 제외합니다."""
    result = {"added": [], "updated": [], "rent": 0, "new_sites": []}
    new_sites = sorted({clean_str(v.get("현장")) for v in vehicles if clean_str(v.get("구분")) == "구매"} - set(site_names()) - {""})
    if new_sites:
        base = len(site_names())
        upsert("sites", [{"현장명": s, "구분": "현장", "표시순서": base + i + 1, "비고": "차량 연동 시 추가"} for i, s in enumerate(new_sites)])
        result["new_sites"] = new_sites
    linked = {a["연결차량번호"]: a for a in load("assets") if a["연결차량번호"]}
    counters = _tag_counters()
    new_recs = []
    for v in vehicles:
        no = clean_str(v.get("차량번호"))
        if clean_str(v.get("구분")) != "구매":
            result["rent"] += 1
            continue
        site = clean_str(v.get("현장")) or "본사"
        status = "수리중" if clean_str(v.get("상태")) == "정비중" else "사용중"
        cost = to_num(v.get("차량 구매 비용"), 0) or 0
        acquired = to_date(v.get("차량 구매일")) or today()
        holder = clean_str(v.get("사용자"))
        if no in linked:
            a = linked[no]
            changes = {"현장": site, "사용자": holder, "품명": clean_str(v.get("차종")) or a["품명"]}
            if a["상태"] in ("사용중", "수리중"):
                changes["상태"] = status
            if not a["취득가액"] and cost:
                changes["취득가액"] = cost
            if any(clean_str(a[k]) != clean_str(changes[k]) for k in changes):
                update_asset(a["태그번호"], changes, actor, "차량 연동", "")
                result["updated"].append(a["태그번호"])
            continue
        tag = _next_tags("차량운반구", 1, acquired.year, counters)[0]
        new_recs.append(_new_asset({"품명": clean_str(v.get("차종")) or "법인차량", "규격": f"차량번호 {no}", "시리얼번호": no,
                                    "계정과목": "차량운반구", "현장": site, "사용자": holder, "취득일자": acquired,
                                    "취득가액": cost, "상태": status, "연결차량번호": no,
                                    "비고": f"법인차량 연동 · 관리자 {clean_str(v.get('관리자'))}".strip()}, tag))
    if new_recs:
        upsert("assets", new_recs)
        add_history_many([(r["태그번호"], "차량 연동", f"법인차량 {r['연결차량번호']} 연동 등록") for r in new_recs], actor)
        result["added"] = [r["태그번호"] for r in new_recs]
    return result


# =========================================================
# 예시 데이터 (보유 약 2만 개 + 처분 약 18만 개)
# =========================================================
SAMPLE_MARK = "[예시]"

# (계정과목, 품명, 규격 후보, 단가 범위)
SAMPLE_CATALOG = [
    ("비품", "노트북", ["LG gram 16", "삼성 갤럭시북4 Pro", "MacBook Pro 14", "Lenovo ThinkPad T14"], (1400000, 2900000)),
    ("비품", "데스크탑 PC", ["삼성 DM500", "LG Z80", "Dell OptiPlex 7010"], (1100000, 1900000)),
    ("비품", "워크스테이션", ["Dell Precision 7960", "HP Z4 G5"], (4500000, 9000000)),
    ("비품", "복합기", ["Canon iR-ADV C5840", "신도리코 D460"], (5000000, 12000000)),
    ("비품", "회의실 프로젝터", ["Epson EB-L630U", "BenQ LU935"], (1500000, 3500000)),
    ("비품", "사무용 책상세트", ["퍼시스 1600", "일룸 모션데스크"], (1000000, 1800000)),
    ("비품", "서버", ["Dell PowerEdge R650", "HPE ProLiant DL380"], (8000000, 18000000)),
    ("비품", "에어컨", ["삼성 무풍 스탠드", "LG 휘센 천장형"], (1800000, 4200000)),
    ("공기구비품", "토탈스테이션", ["Leica TS07", "Topcon GM-50"], (12000000, 22000000)),
    ("공기구비품", "레이저 레벨기", ["Topcon RL-H5A", "Leica Rugby 640"], (2200000, 4200000)),
    ("공기구비품", "핸드파렛트 트럭", ["2.5톤", "3톤"], (1000000, 1600000)),
    ("공기구비품", "오실로스코프", ["Keysight DSOX3054T", "Tektronix MSO44"], (9000000, 18000000)),
    ("공기구비품", "용접기", ["CO2 350A", "TIG 300A"], (1500000, 3800000)),
    ("공기구비품", "드론 측량 장비", ["DJI Matrice 350 RTK"], (10000000, 14000000)),
    ("기계장치", "전동 지게차", ["두산 3톤 BR30S", "현대 2.5톤"], (32000000, 48000000)),
    ("기계장치", "컨베이어 라인", ["벨트형 30m", "롤러형 20m"], (40000000, 90000000)),
    ("기계장치", "발전기", ["100kW 디젤", "60kW 디젤"], (18000000, 32000000)),
    ("기계장치", "공기압축기", ["37kW 스크류", "22kW 스크류"], (12000000, 25000000)),
    ("소프트웨어", "CAD 설계 소프트웨어", ["영구 라이선스", "네트워크 라이선스"], (6000000, 32000000)),
    ("소프트웨어", "ERP 모듈 라이선스", ["회계 모듈", "구매 모듈"], (15000000, 45000000)),
    ("소프트웨어", "오피스 라이선스", ["MS Office LTSC 2024", "한컴오피스 2024"], (400000, 650000)),
    ("차량운반구", "현장 승합차", ["기아 카니발", "현대 스타리아"], (34000000, 42000000)),
    ("차량운반구", "화물차", ["현대 포터2", "기아 봉고3"], (19000000, 24000000)),
    ("소모품", "모니터", ["27인치 QHD", "24인치 FHD", "32인치 4K"], (180000, 520000)),
    ("소모품", "무선 키보드·마우스", ["Logitech MX", "로지텍 MK470"], (40000, 190000)),
    ("소모품", "바코드 스캐너", ["Zebra DS2208", "Honeywell 1472g"], (200000, 380000)),
    ("소모품", "사무용 의자", ["시디즈 T50", "듀오백 Q1"], (180000, 450000)),
    ("소모품", "안전모 세트", ["20개입"], (250000, 400000)),
    ("소모품", "전동드릴", ["보쉬 GSR 18V", "디월트 DCD791"], (150000, 380000)),
    ("소모품", "외장 SSD", ["삼성 T7 1TB", "SanDisk 2TB"], (120000, 260000)),
]
SAMPLE_WEIGHTS = {"비품": 34, "소모품": 32, "공기구비품": 16, "기계장치": 6, "소프트웨어": 7, "차량운반구": 5}
SAMPLE_NAMES = ["김민수", "이지훈", "박서연", "정우성", "한지민", "오세훈", "유나래", "강태호", "서지수", "윤도현",
                "임재현", "최유진", "조성민", "배수아", "노준호", "문가영", "신동욱", "권하늘", "황보람", "송지아"]
SAMPLE_DEPTS = ["경영지원팀", "재무팀", "인사팀", "구매팀", "설계팀", "공무팀", "운영팀", "설비팀", "품질팀", "안전팀", "생산1팀", "생산2팀"]
SAMPLE_VENDORS = {
    "폐기": ["(주)그린리사이클", "한국전자폐기물센터", "에코순환", "대한자원"],
    "매각": ["중고나라 기업몰", "(주)리셀코리아", "서울중고장비", "한빛자산매입"],
    "기증": ["지역아동센터", "OO복지관", "사회적기업 나눔"],
    "사내 경매": ["사내 경매"],
}


def has_sample_data() -> bool:
    return bool(scalar("SELECT 1 FROM assets WHERE \"비고\" LIKE '%[예시]%' LIMIT 1"))


def _shift_months(d: date, months: int) -> date:
    y, m = divmod(d.month - 1 + months, 12)
    return date(d.year + y, m + 1, min(d.day, 28))


def create_sample_data(actor: str = "예시", active_count: int = 20500, closed_count: int = 180000, progress=None) -> dict:
    """보유 자산 약 2만 개 + 처분(폐기·매각·기증·경매·망실) 자산 약 18만 개와 진행 중인 요청·재물조사를 만듭니다.

    모든 예시 자산의 비고에 [예시]가 붙고, 설정에서 한 번에 지울 수 있습니다.
    progress(비율 0~1, 문구)를 넘기면 진행 상황을 알려 줍니다.
    """
    if has_sample_data():
        raise DataError("이미 예시 데이터가 있습니다. 먼저 삭제한 뒤 다시 만드세요.")
    report = progress or (lambda p, t: None)
    rnd = random.Random(2026)
    sites = site_names() or [s["현장명"] for s in DEFAULT_SITES]
    site_w = ([30, 25, 25, 20] + [15] * len(sites))[:len(sites)]
    accounts = list(SAMPLE_WEIGHTS)
    acc_w = [SAMPLE_WEIGHTS[a] for a in accounts]
    catalog = {a: [c for c in SAMPLE_CATALOG if c[0] == a] for a in accounts}
    counters = _tag_counters()
    now = today()
    stamp = now_text()
    sample_actor = f"{SAMPLE_MARK} {actor}"
    blank = {c: "" for c in TABLES["assets"]["cols"]}
    for c in TABLES["assets"]["nums"] + TABLES["assets"]["dates"]:
        blank[c] = None

    def make(acc: str, acquired: date, site: str) -> dict:
        _, name, specs, (lo, hi) = rnd.choice(catalog[acc])
        info = ACCOUNTS[acc]
        rec = dict(blank)
        rec.update({"태그번호": _next_tags(acc, 1, acquired.year, counters)[0], "품명": name, "규격": rnd.choice(specs),
                    "시리얼번호": "" if acc == "소모품" else f"SN{rnd.randint(10000000, 99999999)}",
                    "계정과목": acc, "자산구분": info["cls"], "현장": site, "관리부서": rnd.choice(SAMPLE_DEPTS),
                    "사용자": "" if acc in ("기계장치", "소프트웨어") else rnd.choice(SAMPLE_NAMES),
                    "설치위치": f"{rnd.randint(1, 5)}층" if acc in ("비품", "소모품") else rnd.choice(["자재창고", "현장사무소", "작업장", "1공구", "2공구"]),
                    "취득일자": acquired, "취득가액": round(rnd.uniform(lo, hi), -3), "내용연수": info["life"],
                    "상각방법": info["method"], "잔존가액": 0,
                    "공급업체": rnd.choice(["OO상사", "한빛정보", "대성산업", "미래테크", "에이스장비"]),
                    "상태": "사용중", "등록일시": stamp, "비고": SAMPLE_MARK})
        return rec

    # 1) 보유 자산 (최근 9년 안에 취득, 일부는 내용연수 만료)
    report(0.02, "보유 자산 만드는 중")
    active, hist = [], []
    month_start = date(now.year, now.month, 1)
    for _ in range(active_count):
        acc = rnd.choices(accounts, acc_w)[0]
        rec = make(acc, month_start - timedelta(days=rnd.randint(20, 9 * 365)), rnd.choices(sites, site_w)[0])
        if rnd.random() < 0.055:
            rec["상태"], rec["사용자"] = "유휴", ""
        active.append(rec)
        hist.append({"일시": stamp, "태그번호": rec["태그번호"], "구분": "등록", "내용": "예시 데이터 (보유 자산)", "처리자": sample_actor})

    # 2) 처분 자산 (2004년~ 취득, 내용연수 전후로 처분)
    report(0.15, "처분 자산 만드는 중")
    closed, groups = [], {}
    methods, method_w = ["폐기", "매각", "기증", "사내 경매", "망실"], [62, 25, 6, 4, 3]
    for i in range(closed_count):
        if i % 30000 == 0:
            report(0.15 + 0.3 * i / closed_count, f"처분 자산 만드는 중 ({i:,}/{closed_count:,})")
        acc = rnd.choices(accounts, acc_w)[0]
        life = ACCOUNTS[acc]["life"] or rnd.randint(2, 5)
        acquired = date(rnd.randint(2004, now.year - 1), rnd.randint(1, 12), rnd.randint(1, 28))
        method = rnd.choices(methods, method_w)[0]
        offset = rnd.randint(-6, 36) if method != "망실" else rnd.randint(3 - life * 12, 0)
        disposed = _shift_months(acquired, life * 12 + offset)
        if disposed >= now:
            disposed = now - timedelta(days=rnd.randint(1, 300))
            if disposed <= acquired:
                acquired = disposed - timedelta(days=rnd.randint(200, 2000))
        rec = make(acc, acquired, rnd.choices(sites, site_w)[0])
        cost = rec["취득가액"]
        nbv = cost - accumulated_at(rec, months_used(acquired, disposed))
        proceeds = (round(cost * rnd.uniform(0.02, 0.15), -3) if method == "매각" else
                    round(cost * rnd.uniform(0.03, 0.2), -3) if method == "사내 경매" else 0)
        rec.update({"상태": "망실" if method == "망실" else "처분완료", "처분일자": disposed, "처분방법": method,
                    "처분처": "" if method == "망실" else rnd.choice(SAMPLE_VENDORS[method]),
                    "처분금액": proceeds, "처분손익": round(proceeds - nbv), "사용자": ""})
        closed.append(rec)
        if method != "망실":
            groups.setdefault((disposed.strftime("%y%m"), rec["현장"], method), []).append(rec)
        hist.append({"일시": f"{disposed} 17:00", "태그번호": rec["태그번호"], "구분": "망실 제적" if method == "망실" else "처분 제적",
                     "내용": f"{method}{' · ' + rec['처분처'] if rec['처분처'] else ''} · 처분손익 {rec['처분손익']:,}원 (예시)",
                     "처리자": sample_actor})

    # 3) 처분 요청서 (월·현장·방법별 1건, 제적완료 상태)
    report(0.48, "처분 요청서 만드는 중")
    dsp_rows, seq = [], {}
    for (ym, site, method), items in sorted(groups.items()):
        prefix = f"DS-{ym}-"
        seq[prefix] = seq.get(prefix, 0) + 1
        no = f"{prefix}{seq[prefix]:03d}"
        when = max(a["처분일자"] for a in items)
        proceeds = sum(a["처분금액"] for a in items)
        gl = sum(a["처분손익"] for a in items)
        wipe = needs_secure_wipe(items)
        for a in items:
            a["처분요청번호"] = no
        dsp_rows.append({"요청번호": no, "상태": "제적완료", "요청일시": f"{when - timedelta(days=20)} 09:00",
                         "신청자": f"{SAMPLE_MARK} {site} 담당", "현장": site, "태그번호": ", ".join(a["태그번호"] for a in items),
                         "품명": ", ".join(sorted({a["품명"] for a in items})), "불용사유": "내용연수 만료",
                         "희망 처분방법": method, "상세사유": "정기 불용품 정리 (예시)", "위원회 의견": "불용 판정",
                         "승인일시": f"{when - timedelta(days=10)} 10:00", "처분방법": method, "집행일": when,
                         "처분처": items[0]["처분처"], "처분금액": proceeds, "보안파기 확인": "확인" if wipe else "",
                         "파기증명번호": f"WIPE-{ym}-{rnd.randint(1000, 9999)}" if wipe else "",
                         "장부가 합계": round(proceeds - gl), "처분손익": gl, "제적일시": f"{when} 17:00",
                         "처리이력": f"[{when} 17:00] 장부 제적 · {sample_actor}"})

    report(0.55, f"저장하는 중 (보유 자산 {len(active):,}개)")
    insert_bulk("assets", active, normalized=True)
    report(0.65, f"저장하는 중 (처분 자산 {len(closed):,}개)")
    insert_bulk("assets", closed, normalized=True)
    report(0.8, "처분 요청서·자산 이력 저장하는 중")
    insert_bulk("dsp", dsp_rows)
    insert_bulk("history", hist, normalized=True)

    # 4) 진행 중인 요청·재물조사 (결재 흐름 체험용)
    report(0.9, "진행 중인 요청·재물조사 만드는 중")
    s = (sites + sites)[:4]
    who = lambda name, site: f"{SAMPLE_MARK} {name}({site})"
    adm = f"{SAMPLE_MARK} 예시 관리자"
    create_purchase_request({"현장": s[2], "품명": "GPU 워크스테이션", "규격": "RTX 6000 Ada, 128GB", "계정과목": "비품",
                             "수량": 2, "예상단가": 6800000, "필요일": now + timedelta(days=30),
                             "구매사유": "신규 시뮬레이션 과제 착수. 기존 워크스테이션은 모두 사용 중"}, who("박민수", s[2]))
    pr2 = create_purchase_request({"현장": s[3], "품명": "드론 측량 장비", "규격": "DJI Matrice 350 RTK", "계정과목": "공기구비품",
                                   "수량": 1, "예상단가": 12000000, "구매사유": "토공량 측정 주기 단축"}, who("최동훈", s[3]))
    pr_decide(pr2, True, "예산 범위 내", adm)
    pr_save_quotes(pr2, [{"업체": "에어로텍", "단가": 11850000, "조건": "교육 포함"},
                         {"업체": "스카이맵", "단가": 12300000, "조건": "보증 2년"},
                         {"업체": "드론원", "단가": 11600000, "조건": "납기 3주"}], adm)
    pr_choose_vendor(pr2, 3, "최저가, 납기 허용 범위", adm)
    idle_laptop = next((a for a in active if a["상태"] == "유휴" and a["품명"] == "노트북" and a["현장"] != s[2]), None)
    if idle_laptop:
        create_transfer([idle_laptop["태그번호"]], s[2], "설계팀", "서지수", "신규 입사자 지급용. 유휴 노트북 재배치 요청", who("박민수", s[2]))
    station = next((a for a in active if a["품명"] == "토탈스테이션" and a["상태"] == "사용중"), None)
    if station:
        create_incident(station["태그번호"], "파손", now - timedelta(days=6), "3공구 성토구간", "강태호",
                        "측량 중 삼각대 전도로 본체 낙하. 렌즈부 파손 확인.", "", who("최동훈", station["현장"]))
    old = [get_asset(a["태그번호"]) for a in active if a["품명"] == "노트북" and a["현장"] == s[0] and a["상태"] == "사용중"]
    old = [a for a in old if a and (life_remaining_months(a) or 1) <= 0][:3]
    if old:
        dsp = create_disposal([a["태그번호"] for a in old], "내용연수 만료", "매각", "배터리 노후, 성능 저하", who("김지현", s[0]))
        disposal_decide(dsp[0], True, "중고 매각 후 잔존가치 회수", adm)
    no = start_audit(f"{now.year} 하반기 정기 재물조사", now - timedelta(days=3), "전체", adm)
    items = get_audit_items(no)
    rnd.shuffle(items)
    marks = {r["태그번호"]: {"결과": "미발견" if i % 97 == 5 else "확인", "유휴": i % 61 == 9}
             for i, r in enumerate(items[: int(len(items) * 0.6)])}
    _bulk_audit_mark(no, marks, adm)
    report(1.0, "완료")
    return {"active": len(active), "closed": len(closed), "dsp": len(dsp_rows)}


def _bulk_audit_mark(audit_no: str, marks: dict[str, dict], actor: str):
    stamp = now_text()
    with _lock:
        with _db() as conn:
            conn.executemany('UPDATE audit_items SET "결과"=?, "유휴"=?, "확인자"=?, "확인일시"=? WHERE "회차번호"=? AND "태그번호"=?',
                             [(m["결과"], "Y" if m.get("유휴") else "", actor, stamp, audit_no, t) for t, m in marks.items()])
        s = _audit_counts(audit_no)
        upsert("audits", [dict(lookup("audits", audit_no), 확인=s["found"], 미발견=s["missing"], 유휴=s["idle"])])


def clear_sample_data() -> int:
    """[예시] 자산과 그 자산이 들어간 요청·조사·이력을 지웁니다."""
    like = "%[예시]%"
    with _lock:
        n = scalar('SELECT COUNT(*) FROM assets WHERE "비고" LIKE ?', (like,)) or 0
        conn = _db()
        with conn:
            conn.execute('DELETE FROM history WHERE "처리자" LIKE ? OR "태그번호" IN (SELECT "태그번호" FROM assets WHERE "비고" LIKE ?)', (like, like))
            conn.execute('DELETE FROM audit_items WHERE "회차번호" IN (SELECT "회차번호" FROM audits WHERE "시작자" LIKE ?) '
                         'OR "태그번호" IN (SELECT "태그번호" FROM assets WHERE "비고" LIKE ?)', (like, like))
            conn.execute('DELETE FROM audits WHERE "시작자" LIKE ?', (like,))
            for kind in REQUEST_TABLES:
                conn.execute(f'DELETE FROM {kind} WHERE "처리이력" LIKE ?', (like,))
            conn.execute('DELETE FROM assets WHERE "비고" LIKE ?', (like,))
        for t in ["assets", "audits"] + REQUEST_TABLES:
            _invalidate(t)
        conn.execute("VACUUM")
    return n
