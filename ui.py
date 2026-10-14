"""
공통 화면 요소: CSS, KPI 카드, 배지, 숫자 형식, QR 코드·라벨, 엑셀 내려받기.
법인차량 관리 AI와 같은 카드·배지 스타일을 사용합니다.
"""
from __future__ import annotations

import html
import io
from datetime import date
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl.styles import Font

import data as d

try:
    import qrcode
    from PIL import Image, ImageDraw, ImageFont
except ImportError:  # requirements.txt 설치 전
    qrcode = None
    Image = ImageDraw = ImageFont = None


# =========================================================
# 스타일
# =========================================================
def inject_css():
    st.markdown(
        """
        <style>
        .block-container {padding-top: 1.2rem; padding-bottom: 2rem;}
        .kpi-card {border:1px solid #e6e8eb;border-radius:16px;padding:16px 18px;background:#fff;box-shadow:0 2px 10px rgba(0,0,0,.04);min-height:115px;}
        .kpi-card.warn {border-color:#fcd34d;background:#fffbeb;}
        .kpi-card.crit {border-color:#fca5a5;background:#fef2f2;}
        .kpi-card.ok {border-color:#86efac;background:#f0fdf4;}
        .kpi-title {font-size:14px;color:#64748b;font-weight:700;}
        .kpi-value {font-size:28px;font-weight:800;margin-top:6px;color:#111827;font-variant-numeric:tabular-nums;}
        .kpi-value small {font-size:15px;color:#64748b;font-weight:700;margin-left:2px;}
        .kpi-caption {font-size:12px;color:#94a3b8;margin-top:4px;}
        .badge {display:inline-block;padding:4px 8px;border-radius:999px;font-size:12px;font-weight:700;white-space:nowrap;}
        .approved {background:#dcfce7;color:#166534;}
        .pending {background:#fef3c7;color:#92400e;}
        .rejected {background:#fee2e2;color:#991b1b;}
        .info {background:#dbeafe;color:#1e40af;}
        .muted {background:#f1f5f9;color:#475569;}
        .section-title {font-size:19px;font-weight:800;margin:12px 0 8px;}
        .small-muted {color:#64748b;font-size:12px;}
        .tag {font-family:ui-monospace,SFMono-Regular,Menlo,Consolas,monospace;font-weight:700;background:#f1f5f9;border:1px solid #e2e8f0;border-radius:6px;padding:1px 7px;font-size:13px;}
        .steps {display:flex;flex-wrap:wrap;gap:4px;align-items:center;font-size:11px;margin-top:4px;}
        .steps span {padding:2px 7px;border-radius:6px;background:#f1f5f9;color:#94a3b8;font-weight:700;}
        .steps span.done {background:#dcfce7;color:#166534;}
        .steps span.now {background:#2563eb;color:#fff;}
        .steps span.stop {background:#fee2e2;color:#991b1b;}
        .steps em {color:#cbd5e1;font-style:normal;}
        .todo-card {border:1px solid #e5e7eb;border-radius:14px;padding:14px 16px;background:#fff;}
        .todo-card b {font-size:22px;}
        .asset-head {display:flex;gap:10px;align-items:center;flex-wrap:wrap;margin-bottom:6px;}
        .asset-head .name {font-size:20px;font-weight:800;}
        .ai-box {border:1px solid #bfdbfe;background:#eff6ff;border-radius:14px;padding:12px 16px;}
        </style>
        """,
        unsafe_allow_html=True,
    )


def flash(message: str, kind: str = "success"):
    """다음 화면 갱신 때 보여줄 알림을 남깁니다(대화상자는 rerun 시 닫히므로)."""
    st.session_state.setdefault("_flash", []).append((kind, message))


def show_flash():
    for kind, message in st.session_state.pop("_flash", []):
        getattr(st, kind if kind in ("success", "info", "warning", "error") else "info")(message)


def run_action(fn, *args, success: str | None = None, **kwargs):
    """데이터 작업을 실행하고 오류는 화면에 표시합니다. 성공하면 결과를 돌려줍니다."""
    try:
        result = fn(*args, **kwargs)
    except d.DataError as exc:
        st.error(str(exc))
        return None
    except Exception as exc:  # 예상하지 못한 오류도 화면에서 확인할 수 있게
        st.error(f"처리하지 못했습니다: {exc}")
        return None
    if success:
        flash(success)
    return result if result is not None else True


# =========================================================
# 형식
# =========================================================
def won(value) -> str:
    try:
        return f"{round(float(value)):,}원"
    except (TypeError, ValueError):
        return "-"


def short_won(value) -> str:
    try:
        n = float(value)
    except (TypeError, ValueError):
        return "-"
    if abs(n) >= 1e8:
        return f"{n / 1e8:,.2f}".rstrip("0").rstrip(".") + "억"
    if abs(n) >= 1e4:
        return f"{round(n / 1e4):,}만"
    return f"{round(n):,}"


def fmt_date(value) -> str:
    v = d.to_date(value)
    return v.strftime("%Y-%m-%d") if v else "-"


def kpi(title: str, value, caption: str = "", tone: str = "", unit: str = ""):
    unit_html = f"<small>{html.escape(unit)}</small>" if unit else ""
    st.markdown(
        f"""<div class="kpi-card {tone}"><div class="kpi-title">{html.escape(title)}</div>
        <div class="kpi-value">{html.escape(str(value))}{unit_html}</div>
        <div class="kpi-caption">{html.escape(caption)}</div></div>""",
        unsafe_allow_html=True,
    )


STATUS_TONE = {
    "사용중": "approved", "유휴": "pending", "수리중": "info", "처분대기": "rejected", "망실": "rejected", "처분완료": "muted",
    "요청": "pending", "신청": "pending", "접수": "pending", "승인": "info", "발주": "info", "검수완료": "info",
    "조사완료": "info", "집행완료": "info", "등록완료": "approved", "완료": "approved", "종결": "approved",
    "제적완료": "approved", "반려": "rejected", "진행중": "pending",
}


def badge(status: str) -> str:
    return f'<span class="badge {STATUS_TONE.get(status, "muted")}">{html.escape(status or "-")}</span>'


def steps_html(kind: str, status: str) -> str:
    flow = d.FLOWS[kind]
    if status == d.REJECTED:
        return f'<div class="steps"><span class="done">{flow[0]}</span><em>›</em><span class="stop">반려</span></div>'
    idx = flow.index(status) if status in flow else 0
    parts = []
    for i, s in enumerate(flow):
        cls = "done" if i < idx else "now" if i == idx else ""
        if i == len(flow) - 1 and i == idx:
            cls = "done"
        parts.append(f'<span class="{cls}">{s}</span>')
    return '<div class="steps">' + "<em>›</em>".join(parts) + "</div>"


# =========================================================
# 표
# =========================================================
ASSET_VIEW_COLS = ["태그번호", "품명", "규격", "계정과목", "현장", "관리부서", "사용자", "설치위치", "취득일자", "취득가액", "장부가액", "상태"]


def asset_frame(assets: list[dict], extra: list[str] | None = None) -> pd.DataFrame:
    rows = []
    for a in assets:
        left = d.life_remaining_months(a)
        row = {c: a.get(c, "") for c in ASSET_VIEW_COLS if c != "장부가액"}
        row["취득일자"] = d.to_date(a["취득일자"])
        row["취득가액"] = d.to_num(a["취득가액"], 0)
        row["장부가액"] = round(d.book_value(a))
        row["내용연수"] = "만료" if left is not None and left <= 0 and d.is_active(a) else (f"{left}개월 남음" if left is not None and d.is_active(a) else "-")
        for c in extra or []:
            row[c] = a.get(c, "")
        rows.append(row)
    cols = ASSET_VIEW_COLS + ["내용연수"] + (extra or [])
    return pd.DataFrame(rows, columns=cols)


ASSET_COLUMN_CONFIG = {
    "취득가액": st.column_config.NumberColumn("취득가액", format="localized"),
    "장부가액": st.column_config.NumberColumn("장부가액", format="localized"),
    "취득일자": st.column_config.DateColumn("취득일자", format="YYYY-MM-DD"),
}


def money_config(*cols) -> dict:
    return {c: st.column_config.NumberColumn(c, format="localized") for c in cols}


def paginate(rows: list, key: str, page_size: int = 1000) -> list:
    """긴 목록을 페이지로 나눠 현재 페이지 행만 돌려줍니다(수십만 행도 화면이 느려지지 않게)."""
    total = len(rows)
    pages = max(1, -(-total // page_size))
    if st.session_state.get(key, 1) > pages:
        st.session_state[key] = 1
    if pages == 1:
        return rows
    c1, c2 = st.columns([1, 4])
    page = c1.number_input("페이지", min_value=1, max_value=pages, step=1, key=key)
    start = (page - 1) * page_size
    end = min(total, start + page_size)
    c2.markdown(f"<div style='padding-top:2rem' class='small-muted'>{start + 1:,}–{end:,} / 전체 {total:,}개 · {pages:,}쪽</div>",
                unsafe_allow_html=True)
    return rows[start:end]


def export_panel(key: str, filename: str, build, row_count: int, xlsx_limit: int = 300_000):
    """버튼을 눌렀을 때만 파일을 만듭니다. build() → {시트명: (열 목록, 행 목록)}.

    수십만 행을 매 화면마다 엑셀로 만들면 느려지므로, 만들기 → 내려받기 두 단계로 나눴습니다.
    """
    c1, c2, c3 = st.columns([1, 1, 2])
    if c1.button(f"📗 엑셀 만들기 ({row_count:,}행)", key=f"{key}-mk-x", width="stretch",
                 disabled=row_count == 0 or row_count > xlsx_limit,
                 help="행이 많으면 수십 초 걸릴 수 있습니다" if row_count < xlsx_limit else "행이 너무 많아 CSV로 내보내세요"):
        with st.spinner(f"엑셀 파일을 만드는 중입니다 ({row_count:,}행)..."):
            st.session_state[f"{key}-file"] = ("xlsx", d.rows_to_xlsx(build()))
    if c2.button("📄 CSV 만들기", key=f"{key}-mk-c", width="stretch", disabled=row_count == 0,
                 help="엑셀에서 바로 열립니다. 행이 많을 때 가장 빠릅니다"):
        with st.spinner("CSV 파일을 만드는 중입니다..."):
            sheets = build()
            cols, rows = next(iter(sheets.values()))
            st.session_state[f"{key}-file"] = ("csv", d.rows_to_csv(cols, rows))
    made = st.session_state.get(f"{key}-file")
    if made:
        kind, blob = made
        with c3:
            st.download_button(f"⬇️ {filename}.{kind} 내려받기 ({len(blob) / 1e6:.1f}MB)", blob, file_name=f"{filename}.{kind}",
                               mime="text/csv" if kind == "csv" else None, key=f"{key}-dl", width="stretch", type="primary")


def excel_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
            ws = writer.sheets[name[:31]]
            for idx, col in enumerate(df.columns, start=1):
                letter = ws.cell(row=1, column=idx).column_letter
                ws.column_dimensions[letter].width = max(10, min(40, len(str(col)) * 2 + 4))
                cell = ws.cell(row=1, column=idx)
                cell.font = Font(bold=True)
                if str(col) in d.MONEY_COLS or str(col).endswith(("금액", "가액", "상각비", "장부가", "취득")):
                    for r in range(2, len(df) + 2):
                        ws.cell(row=r, column=idx).number_format = "#,##0"
            ws.freeze_panes = "B2"
    return buf.getvalue()


# =========================================================
# QR 코드 / 라벨
# =========================================================
FONT_CANDIDATES = [
    "C:/Windows/Fonts/malgunbd.ttf", "C:/Windows/Fonts/malgun.ttf",
    "/System/Library/Fonts/AppleSDGothicNeo.ttc", "/Library/Fonts/AppleGothic.ttf",
    "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf", "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
]


def _font(size: int, bold: bool = False):
    order = FONT_CANDIDATES if bold else [f for f in FONT_CANDIDATES if "bd" not in f.lower() and "bold" not in f.lower()] + FONT_CANDIDATES
    for path in order:
        if Path(path).exists():
            try:
                return ImageFont.truetype(path, size)
            except OSError:
                continue
    return ImageFont.load_default(size=size)


def korean_font_available() -> bool:
    return any(Path(p).exists() for p in FONT_CANDIDATES)


def qr_payload(tag: str) -> str:
    base = d.clean_str(d.get_settings().get("QR 기본주소")).rstrip("/")
    return f"{base}/?tag={tag}" if base else tag


def qr_image(text: str, box_size: int = 10):
    if qrcode is None:
        raise d.DataError("qrcode 패키지가 없습니다. 터미널에서 pip install -r requirements.txt 를 실행하세요.")
    qr = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=box_size, border=2)
    qr.add_data(text)
    qr.make(fit=True)
    return qr.make_image(fill_color="black", back_color="white").convert("RGB")


def qr_png(text: str, box_size: int = 10) -> bytes:
    buf = io.BytesIO()
    qr_image(text, box_size).save(buf, format="PNG")
    return buf.getvalue()


def _fit(draw, text: str, font, max_w: int) -> str:
    if draw.textlength(text, font=font) <= max_w:
        return text
    while text and draw.textlength(text + "…", font=font) > max_w:
        text = text[:-1]
    return text + "…"


def _draw_label(img, asset: dict, x: int, y: int, w: int, h: int, company: str):
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle([x, y, x + w, y + h], radius=14, outline=(40, 40, 40), width=3, fill="white")
    pad = 18
    qr_size = h - pad * 2
    qr = qr_image(qr_payload(asset["태그번호"]), box_size=10).resize((qr_size, qr_size), Image.NEAREST)
    img.paste(qr, (x + pad, y + pad))
    tx = x + pad * 2 + qr_size
    tw = w - (tx - x) - pad
    f_small, f_tag, f_name, f_meta = _font(22), _font(42, True), _font(30, True), _font(24)
    draw.text((tx, y + pad), _fit(draw, (company + " " if company else "") + "자산관리표", f_small, tw), font=f_small, fill=(90, 90, 90))
    draw.text((tx, y + pad + 34), asset["태그번호"], font=f_tag, fill="black")
    draw.text((tx, y + pad + 94), _fit(draw, asset["품명"], f_name, tw), font=f_name, fill="black")
    site = asset["현장"] + (f" · {asset['관리부서']}" if asset["관리부서"] else "")
    draw.text((tx, y + pad + 140), _fit(draw, site, f_meta, tw), font=f_meta, fill=(50, 50, 50))
    draw.text((tx, y + pad + 176), _fit(draw, f"취득 {fmt_date(asset['취득일자'])} · {asset['계정과목']}", f_meta, tw), font=f_meta, fill=(50, 50, 50))


def label_png(asset: dict) -> bytes:
    img = Image.new("RGB", (900, 300), "white")
    _draw_label(img, asset, 4, 4, 892, 292, d.clean_str(d.get_settings().get("회사명")))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def labels_pdf(assets: list[dict]) -> bytes:
    """A4(200dpi) 한 장에 2열 × 7행 라벨. 여러 장이면 PDF 여러 페이지."""
    company = d.clean_str(d.get_settings().get("회사명"))
    page_w, page_h, margin, gap = 1654, 2339, 70, 24
    cols, rows = 2, 7
    lw = (page_w - margin * 2 - gap * (cols - 1)) // cols
    lh = (page_h - margin * 2 - gap * (rows - 1)) // rows
    pages = []
    per = cols * rows
    for p in range(0, len(assets), per):
        img = Image.new("RGB", (page_w, page_h), "white")
        for i, a in enumerate(assets[p:p + per]):
            c, r = i % cols, i // cols
            _draw_label(img, a, margin + c * (lw + gap), margin + r * (lh + gap), lw, lh, company)
        pages.append(img)
    if not pages:
        return b""
    buf = io.BytesIO()
    pages[0].save(buf, format="PDF", save_all=True, append_images=pages[1:], resolution=200)
    return buf.getvalue()


LABEL_MAX = 500  # 한 번에 만드는 라벨 수 상한 (A4 약 36장)


def label_download(assets: list[dict], key: str, label: str = "🏷️ QR 라벨 PDF 내려받기"):
    if not assets:
        return
    if len(assets) > LABEL_MAX:
        st.warning(f"라벨은 한 번에 {LABEL_MAX}개까지 만들 수 있습니다. 표에서 출력할 자산을 고르거나 필터로 줄여 주세요 (현재 {len(assets):,}개).")
        return
    try:
        pdf = labels_pdf(assets)
    except d.DataError as exc:
        st.error(str(exc))
        return
    name = f"자산라벨_{assets[0]['태그번호']}.pdf" if len(assets) == 1 else f"자산라벨_{date.today():%Y%m%d}_{len(assets)}개.pdf"
    st.download_button(label, pdf, file_name=name, mime="application/pdf", key=key, width="stretch")
    if not korean_font_available():
        st.caption("※ 한글 글꼴을 찾지 못해 라벨의 한글이 깨질 수 있습니다. 맑은 고딕(Windows)이 있으면 자동으로 사용합니다.")
